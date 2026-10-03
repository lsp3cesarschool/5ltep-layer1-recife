"""Reading a published table as a stream: no raw data is ever stored or committed.

The format is decided by the file's *content*, not by the format the portal declares:

  CSV, JSON, XML   read straight from the HTTP response, record by record, while the SHA-256 is
                   computed; nothing goes to disk;
  ZIP              downloaded to a temporary file (its directory is at the end); every member that
                   is a table is read (CSV/TXT, JSON, XML, spreadsheets, Parquet, and zips inside
                   the zip); a zip with no table inside is not tabular data. A zip larger than the
                   runner's disk is read as it streams instead, member by member (local headers);
  XLSX, XLS, ODS   spreadsheets, downloaded to a temporary file; each sheet is a table, its rows read
                   one at a time (XLSX and ODS) or a sheet at a time (XLS, at most 65,536 rows);
  Parquet          downloaded to a temporary file and read in batches.

Every format becomes the same thing for the rest of the pipeline: a Table whose rows are lists of
strings, header first. JSON and XML are read as lists of records (an array of objects; repeated
elements under the root); a nested value becomes its JSON text. Encoding and delimiter of text
formats are detected from the first 64 KB; bytes that do not decode later are counted, not fatal.
"""

import codecs
import contextlib
import csv
import datetime as dt
import hashlib
import io
import json
import logging
import sys
import tempfile
import time
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlparse

import requests

from src import ckan, config

logger = logging.getLogger(__name__)

csv.field_size_limit(min(sys.maxsize, 2**31 - 1))   # geometry columns hold very long cells

PEEK = 1 << 16
HEADER_SAMPLE = 1000          # records of JSON/XML read to collect the column names
_decode_errors = [0]


def _count_and_replace(exc: UnicodeDecodeError):
    _decode_errors[0] += 1
    return "�", exc.end


codecs.register_error("l1count", _count_and_replace)


class NotTabular(Exception):
    """The resource holds no table. `kind`: "no-tables" (a container of documents or maps: not tabular
    data, left out of the tabular universe), "html" or "pdf" (the link returns a page or a document
    instead of the table: a failure of the link), "too-large"."""

    def __init__(self, message: str, kind: str = "no-tables"):
        super().__init__(message)
        self.kind = kind


class HashingReader(io.RawIOBase):
    """Wraps a byte stream, computing SHA-256 and size of everything read through it."""

    def __init__(self, raw):
        self.raw, self.sha, self.size = raw, hashlib.sha256(), 0

    def readable(self) -> bool:
        return True

    def readinto(self, b) -> int:
        data = self.raw.read(len(b))
        n = len(data)
        b[:n] = data
        self.sha.update(data)
        self.size += n
        return n


def detect_encoding(sample: bytes) -> str:
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    cut = sample[:sample.rfind(b"\n") + 1] or sample
    try:
        cut.decode("utf-8")
        return "utf-8-sig" if sample.startswith(codecs.BOM_UTF8) else "utf-8"
    except UnicodeDecodeError as exc:
        # A multi-byte character cut at the end of the sample is not an error.
        if exc.start >= len(cut) - 3:
            return "utf-8"
        return "cp1252"


def detect_delimiter(text: str) -> str:
    lines = text.splitlines()[:50]
    try:
        return csv.Sniffer().sniff("\n".join(lines), delimiters=";,\t|").delimiter
    except csv.Error:
        first = lines[0] if lines else ""
        return max(";,\t|", key=first.count)


@dataclass
class Table:
    """One table of a resource: the resource itself, a member of its zip, or a sheet."""
    member: str | None
    encoding: str | None          # text formats only
    delimiter: str | None         # CSV only
    rows: object                  # iterator of lists of strings, header first
    decode_errors: list = field(default_factory=lambda: _decode_errors)
    fmt: str = "csv"


def cell(value) -> str:
    """Any value of any format as the text a CSV would hold."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (int, Decimal)):
        return str(value)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def sniff(sample: bytes) -> str:
    """csv, json, xml, zip, parquet or xls, from the first bytes; NotTabular for pages and documents."""
    if sample.startswith(b"PK\x03\x04"):
        return "zip"
    if sample.startswith(b"PAR1"):
        return "parquet"
    if sample.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        text = sample.decode("utf-16", errors="ignore")
    else:
        text = sample.decode("utf-8", errors="ignore")
    head = text.lstrip("﻿ \r\n\t").lower()
    if head.startswith(("<!doctype html", "<html")):
        raise NotTabular("the link returns an HTML page, not a table", "html")
    if head.startswith("%pdf"):
        raise NotTabular("the link returns a PDF document, not a table", "pdf")
    if head.startswith("<"):
        return "xml"
    if head.startswith(("[", "{")):
        return "json"
    return "csv"


# --- text formats, streamed -----------------------------------------------------------

def _csv_table(buffered, member: str | None) -> Table:
    sample = buffered.peek(PEEK)[:PEEK]
    encoding = detect_encoding(sample)
    delimiter = detect_delimiter(sample.decode(encoding, errors="replace"))
    text = io.TextIOWrapper(buffered, encoding=encoding, errors="l1count", newline="")
    return Table(member, encoding, delimiter, csv.reader(text, delimiter=delimiter), fmt="csv")


def _records_to_rows(records):
    """Lists of strings, header first, from an iterator of dicts. The header is the union of the
    keys of the first HEADER_SAMPLE records; a later new key is left out (and the row counts as
    ragged, see validate.py)."""
    first, header, seen = [], [], set()
    for rec in records:
        first.append(rec)
        for k in rec:
            if k not in seen:
                seen.add(k)
                header.append(k)
        if len(first) >= HEADER_SAMPLE:
            break
    yield [str(h) for h in header]
    for rec in first:
        yield _row(rec, header)
    for rec in records:
        yield _row(rec, header)


def _row(rec: dict, header: list) -> list[str]:
    row = [cell(rec.get(h)) for h in header]
    if any(k not in header for k in rec):
        row.append("")                        # a key outside the header: the row is ragged
    return row


def _json_records(stream):
    """Records of a JSON document: the top-level array, or the first array under a top-level key."""
    import ijson

    target, builder, depth = None, None, 0
    for prefix, event, value in ijson.parse(stream):
        if builder is not None:
            builder.event(event, value)
            depth += event in ("start_map", "start_array")
            depth -= event in ("end_map", "end_array")
            if depth == 0:
                yield builder.value if isinstance(builder.value, dict) else {"value": builder.value}
                builder = None
            continue
        if target is None:
            if prefix == "" and event == "start_array":
                target = "item"
            elif event == "start_array" and prefix and "." not in prefix:
                target = prefix + ".item"
            continue
        if prefix == target and event in ("start_map", "start_array"):
            builder, depth = ijson.ObjectBuilder(), 1
            builder.event(event, value)
        elif prefix == target and event in ("string", "number", "boolean", "null"):
            yield {"value": value}


def _json_table(buffered, member: str | None) -> Table:
    sample = buffered.peek(PEEK)[:PEEK]
    encoding = detect_encoding(sample)
    if encoding == "utf-16":                  # ANEEL publishes JSON in UTF-16: transcode for the parser
        stream = io.TextIOWrapper(buffered, encoding="utf-16", errors="l1count")
    else:
        stream = buffered
    return Table(member, encoding, None, _records_to_rows(_json_records(stream)), fmt="json")


def _xml_records(stream):
    """Records of an XML document: the elements directly under the root; fields are their child
    elements (or, when there are none, their attributes)."""
    from defusedxml.ElementTree import iterparse

    depth = 0
    for event, elem in iterparse(stream, events=("start", "end")):
        if event == "start":
            depth += 1
            continue
        depth -= 1
        if depth == 1:
            children = list(elem)
            rec = {c.tag.split("}")[-1]: (c.text or "").strip() for c in children} if children else dict(elem.attrib)
            yield rec
            elem.clear()


def _xml_table(buffered, member: str | None) -> Table:
    encoding = detect_encoding(buffered.peek(PEEK)[:PEEK])
    return Table(member, encoding, None, _records_to_rows(_xml_records(buffered)), fmt="xml")


def _stream_table(buffered, kind: str, member: str | None) -> Table:
    return {"csv": _csv_table, "json": _json_table, "xml": _xml_table}[kind](buffered, member)


# --- formats read from a file on disk ---------------------------------------------------

def _xlsx_tables(path: Path, member: str | None):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = ([cell(v) for v in r] for r in ws.iter_rows(values_only=True))
            yield Table(_name(member, ws.title), None, None, rows, fmt="xlsx")
    finally:
        wb.close()


def _xls_tables(path: Path, member: str | None):
    import xlrd

    # The format holds at most 65,536 rows x 256 columns per sheet: a sheet always fits in memory.
    # Files saved by tools other than Excel often trip xlrd's strict check of the compound document
    # ("Workbook corruption: seen[2] == 4") while their cells read fine.
    book = xlrd.open_workbook(str(path), on_demand=True, ignore_workbook_corruption=True)
    for sheet in book.sheets():
        rows = ([cell(v) for v in sheet.row_values(i)] for i in range(sheet.nrows))
        yield Table(_name(member, sheet.name), None, None, rows, fmt="xls")


ODF_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
ODF_OFFICE = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
ODF_TEXT_P = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}p"


def _ods_value(c) -> str:
    kind = c.get(f"{ODF_OFFICE}value-type")
    if kind in ("float", "percentage", "currency") and c.get(f"{ODF_OFFICE}value") not in (None, ""):
        return cell(float(c.get(f"{ODF_OFFICE}value")))
    if kind == "boolean":
        return (c.get(f"{ODF_OFFICE}boolean-value") or "").lower()
    if kind in ("date", "time"):
        return c.get(f"{ODF_OFFICE}{kind}-value") or ""
    return "\n".join("".join(p.itertext()) for p in c if p.tag == ODF_TEXT_P)


def _ods_tables(path: Path, member: str | None):
    """Each sheet of an ODS, its rows read one at a time from content.xml (never the whole sheet in
    memory). Repeated empty cells and rows (how ODS fills a sheet to its edge) are not materialised;
    rows are padded to the header's width with empty cells, as a CSV export would."""
    from defusedxml.ElementTree import iterparse

    with zipfile.ZipFile(path) as zf, zf.open("content.xml") as fh:
        events = iterparse(fh, events=("start", "end"))
        stack: list = []

        def rows_of(table_el):
            width, empty_rows = None, 0
            for ev, el in events:
                if ev == "start":
                    stack.append(el)
                    continue
                stack.pop()
                if el is table_el:
                    return                                  # trailing empty rows are dropped
                if el.tag != f"{ODF_TABLE}table-row":
                    continue
                values, empty_cells = [], 0
                for c in el:
                    if c.tag in (f"{ODF_TABLE}table-cell", f"{ODF_TABLE}covered-table-cell"):
                        n = int(c.get(f"{ODF_TABLE}number-columns-repeated") or 1)
                        v = _ods_value(c)
                        if v == "":
                            empty_cells += n
                        else:
                            values += [""] * empty_cells + [v] * n
                            empty_cells = 0
                repeat = int(el.get(f"{ODF_TABLE}number-rows-repeated") or 1)
                if stack:
                    stack[-1].remove(el)                    # free the row once read
                if not values:
                    empty_rows += repeat
                    continue
                if width is None:
                    width = len(values)
                for _ in range(empty_rows):
                    yield [""] * width
                empty_rows = 0
                values += [""] * (width - len(values))
                for _ in range(repeat):
                    yield list(values)

        for ev, el in events:
            if ev == "end":
                stack.pop()
                continue
            stack.append(el)
            if el.tag == f"{ODF_TABLE}table":
                rows = rows_of(el)
                yield Table(_name(member, el.get(f"{ODF_TABLE}name") or ""), None, None, rows, fmt="ods")
                for _ in rows:                              # rows the caller did not read
                    pass
                if stack:                                   # the sheet ended: free it
                    stack[-1].remove(el)


def _parquet_tables(path: Path, member: str | None):
    import pyarrow.parquet as pq

    pf = pq.ParquetFile(path)
    names = pf.schema_arrow.names

    def rows():
        yield list(names)
        for batch in pf.iter_batches(batch_size=20000):
            cols = [batch.column(i).to_pylist() for i in range(batch.num_columns)]
            for values in zip(*cols):
                yield [cell(v) for v in values]

    yield Table(member, None, None, rows(), fmt="parquet")


def _name(member: str | None, sheet: str) -> str:
    return f"{member}#{sheet}" if member else sheet


def _container_kind(path: Path) -> str:
    """zip, xlsx or ods for a file that starts with PK."""
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "mimetype" in names and zf.read("mimetype").startswith(b"application/vnd.oasis.opendocument.spreadsheet"):
            return "ods"
    return "zip"


def _file_tables(path: Path, kind: str, member: str | None, tmp: str, depth: int, digest: dict):
    if kind == "zip":
        kind = _container_kind(path)
        if member is None:
            digest["kind"] = kind
    if kind in ("xlsx", "ods", "xls", "parquet") and path.suffix != f".{kind}":
        path = path.replace(path.with_name(f"{path.name}.{kind}"))   # some readers check the extension
    if kind == "xlsx":
        yield from _xlsx_tables(path, member)
    elif kind == "ods":
        yield from _ods_tables(path, member)
    elif kind == "xls":
        yield from _xls_tables(path, member)
    elif kind == "parquet":
        yield from _parquet_tables(path, member)
    else:
        yield from _zip_tables(path, member, tmp, depth, digest)


TABLE_EXTENSIONS = (".csv", ".txt", ".tsv", ".json", ".xml", ".xlsx", ".xls", ".ods", ".parquet", ".zip")


def _zip_tables(path: Path, outer: str | None, tmp: str, depth: int, digest: dict):
    with zipfile.ZipFile(path) as zf:
        infos = [i for i in zf.infolist() if not i.is_dir() and "__MACOSX/" not in i.filename
                 and i.filename.lower().endswith(TABLE_EXTENSIONS)]
        if depth >= config.ZIP_MAX_DEPTH:                   # zips inside zips, up to a few levels
            infos = [i for i in infos if not i.filename.lower().endswith(".zip")]
        found = 0
        for info in sorted(infos, key=lambda i: i.filename):
            name = f"{outer}/{info.filename}" if outer else info.filename
            with zf.open(info) as raw:
                buffered = io.BufferedReader(raw, PEEK * 4)
                try:
                    kind = sniff(buffered.peek(PEEK)[:PEEK])
                except NotTabular:
                    continue
                if kind in ("csv", "json", "xml"):
                    found += 1
                    yield _stream_table(buffered, kind, name)
                    continue
                inner = Path(tmp) / f"member-{depth}-{found}"
                with open(inner, "wb") as fh:
                    while chunk := buffered.read(1 << 20):
                        fh.write(chunk)
            for table in _file_tables(inner, kind, name, tmp, depth + 1, digest):
                found += 1
                yield table
            inner.unlink(missing_ok=True)
        if outer is None:
            digest["members"] = found
        if not found and outer is None:
            names = [i.filename for i in zf.infolist()][:5]
            raise NotTabular(f"zip without tables (first entries: {names})", "no-tables")


class _ChunkReader(io.RawIOBase):
    """A file-like object over an iterator of byte chunks."""

    def __init__(self, chunks):
        self.it, self.buf = iter(chunks), b""

    def readable(self) -> bool:
        return True

    def readinto(self, b) -> int:
        while not self.buf:
            try:
                self.buf = next(self.it)
            except StopIteration:
                return 0
        n = min(len(b), len(self.buf))
        b[:n], self.buf = self.buf[:n], self.buf[n:]
        return n


def _drain(stream) -> None:
    try:
        while stream.read(1 << 20):
            pass
    except ValueError:
        pass                    # a text wrapper that read to the end closed it


def _first_member(sample: bytes) -> str:
    """Name of the first entry of a zip, from its first local header."""
    if not sample.startswith(b"PK\x03\x04") or len(sample) < 30:
        return ""
    n = int.from_bytes(sample[26:28], "little")
    return sample[30:30 + n].decode("utf-8", errors="replace")


def _is_office(first_member: str) -> bool:
    """An XLSX/ODS (or another office document) rather than a plain zip of files."""
    return first_member in ("mimetype", "[Content_Types].xml") \
        or first_member.startswith(("xl/", "_rels/", "docProps/", "META-INF/"))


def _zip_stream_tables(stream, outer: str | None, tmp: str, depth: int, digest: dict):
    """A zip read as it streams, member by member, from the local header of each: for a zip larger
    than the runner's disk. Members are read in the order they are stored."""
    from stream_unzip import stream_unzip

    found = 0
    for i, (raw_name, _size, chunks) in enumerate(stream_unzip(iter(lambda: stream.read(1 << 16), b""))):
        try:
            entry = raw_name.decode("utf-8")
        except UnicodeDecodeError:
            entry = raw_name.decode("cp437")
        if i == 0 and outer is None and _is_office(entry):
            raise NotTabular("a spreadsheet larger than the runner's disk", "too-large")
        name = f"{outer}/{entry}" if outer else entry
        member = io.BufferedReader(_ChunkReader(chunks), PEEK * 4)
        stopped = False
        try:
            lower = entry.lower()
            if entry.endswith("/") or "__MACOSX/" in entry or not lower.endswith(TABLE_EXTENSIONS) \
                    or (lower.endswith(".zip") and depth >= config.ZIP_MAX_DEPTH):
                continue
            try:
                kind = sniff(member.peek(PEEK)[:PEEK])
            except NotTabular:
                continue
            if kind in ("csv", "json", "xml"):
                found += 1
                yield _stream_table(member, kind, name)
                continue
            if kind == "zip" and not _is_office(_first_member(member.peek(PEEK))):
                for table in _zip_stream_tables(member, name, tmp, depth + 1, digest):
                    found += 1
                    yield table
                continue
            inner = Path(tmp) / f"member-{depth}-{found}"
            with open(inner, "wb") as fh:
                while chunk := member.read(1 << 20):
                    fh.write(chunk)
            for table in _file_tables(inner, kind, name, tmp, depth + 1, digest):
                found += 1
                yield table
            inner.unlink(missing_ok=True)
        except GeneratorExit:
            stopped = True      # the caller is done (only a header was read): nothing more is downloaded
            raise
        finally:
            if not stopped:
                _drain(member)  # each member must be read to its end before the next one
    if outer is None:
        digest["members"] = found
        if not found:
            raise NotTabular("zip without tables", "no-tables")


def _streamed_zip(buffered, hashing, tmp: str, digest: dict):
    digest["streamed"] = True
    yield from _zip_stream_tables(buffered, None, tmp, 0, digest)
    _drain(buffered)            # the central directory, so that the hash covers every byte
    digest.update(sha256=hashing.sha.hexdigest(), bytes=hashing.size, kind="zip")


def _serve(source):
    """Hand the source over; when the caller is done (even halfway, as when only the header is read),
    close its readers before the temporary files go away."""
    try:
        yield source
    finally:
        source.tables.close()


class _DiskFull(Exception):
    pass


# --- the resource ---------------------------------------------------------------------

@dataclass
class Source:
    tables: list            # generator of Table
    digest: dict            # filled when the whole file has been read: sha256, bytes, kind


DISK_KINDS = ("zip", "parquet", "xls")


@contextlib.contextmanager
def open_resource(url: str, max_disk_bytes: int | None = None, stream_zip: bool = False):
    """Yields a Source; `digest` is complete once every table has been read to the end. A format that
    must go to disk (zip, spreadsheets, Parquet) is downloaded to a temporary file when it fits in
    `max_disk_bytes` (the runner's disk); a zip that does not is read as it streams; a Parquet or a
    spreadsheet that does not cannot be read on this machine ("too-large"). `stream_zip`: a zip is
    always read as it streams, so that a caller that stops early stops the download too."""
    limit = max_disk_bytes or config.MAX_ZIP_BYTES
    _decode_errors[0] = 0
    t0 = time.monotonic()
    resp = ckan.get(url, stream=True, timeout=config.HTTP_TIMEOUT_S)
    # Network figures: time until the server answered (connection, retries, first byte) and, for a
    # file read from disk, the download alone (a text format is validated while it streams).
    timing = {"first_byte_s": round(time.monotonic() - t0, 2)}
    resp.raw.decode_content = True
    hashing = HashingReader(resp.raw)
    buffered = io.BufferedReader(hashing, PEEK * 4)
    h = resp.headers
    # Change signals of the server, compared with a HEAD request in later runs (work.py).
    digest: dict = {"timing": timing,
                    "http": {"etag": h.get("ETag"), "last_modified": h.get("Last-Modified"),
                             "content_length": int(h["Content-Length"]) if (h.get("Content-Length") or "").isdigit() else None,
                             # how the file was delivered (the dashboard's file delivery findings)
                             "content_type": ckan.media_type(h.get("Content-Type")),
                             "served_by": urlparse(getattr(resp, "url", None) or url).hostname}}
    try:
        sample = buffered.peek(PEEK)[:PEEK]
        kind = sniff(sample)
        length = digest["http"]["content_length"] or 0
        streamable = kind == "zip" and not _is_office(_first_member(sample))
        if kind in DISK_KINDS and length > limit and not streamable:
            raise NotTabular(f"file of {length / 1e9:.1f} GB is larger than the runner's disk", "too-large")
        if streamable and (length > limit or stream_zip):
            with tempfile.TemporaryDirectory() as tmp:
                yield from _serve(Source(_streamed_zip(buffered, hashing, tmp, digest), digest))
        elif kind in DISK_KINDS:
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "resource"
                try:
                    with open(path, "wb") as fh:
                        while chunk := buffered.read(1 << 20):
                            fh.write(chunk)
                            if hashing.size > limit:
                                raise _DiskFull
                except _DiskFull:
                    if not streamable:
                        raise NotTabular("file larger than the runner's disk", "too-large")
                    path.unlink(missing_ok=True)
                    # the server gave no size and the zip does not fit: read it again, as a stream
                    resp.close()
                    resp = ckan.get(url, stream=True, timeout=config.HTTP_TIMEOUT_S)
                    resp.raw.decode_content = True
                    hashing = HashingReader(resp.raw)
                    yield from _serve(Source(_streamed_zip(io.BufferedReader(hashing, PEEK * 4), hashing, tmp, digest), digest))
                else:
                    timing["download_s"] = round(time.monotonic() - t0, 2)
                    digest.update(sha256=hashing.sha.hexdigest(), bytes=hashing.size, kind=kind)
                    yield from _serve(Source(_file_tables(path, kind, None, tmp, 0, digest), digest))
        else:
            def single():
                yield _stream_table(buffered, kind, None)
                # read whatever the reader left (e.g. a trailing newline) so the hash is complete
                try:
                    while buffered.read(1 << 20):
                        pass
                except ValueError:
                    pass       # a text wrapper that read to the end closed it: every byte was hashed
                digest.update(sha256=hashing.sha.hexdigest(), bytes=hashing.size, kind=kind)
            yield from _serve(Source(single(), digest))
    finally:
        resp.close()


def peek_header(url: str, rows: int = 50, max_disk_bytes: int | None = None,
                stream_zip: bool = False) -> tuple[str, list[str]]:
    """(format, header) of a resource, reading as little as the format allows (another format of a
    table already validated in full: only its columns are compared; a file never validated: its columns
    link it to its dictionary).

    `stream_zip`: a zip is read from its start, member by member (each one has its own local header),
    and the download stops after the first table's header: no byte range is needed (Recife's file
    server ignores them), and a 1 GB zip costs a few kilobytes. The members come in the order they are
    stored, not by name as in a validation; a zip that cannot be read so (an unusual layout) is
    downloaded whole."""
    if stream_zip:
        from stream_unzip import UnzipError
        try:
            return _peek(url, rows, max_disk_bytes, True)
        except UnzipError as exc:
            logger.info("%s could not be read as a stream (%s); downloading it whole", url, exc)
    return _peek(url, rows, max_disk_bytes, False)


def _peek(url: str, rows: int, max_disk_bytes: int | None, stream_zip: bool) -> tuple[str, list[str]]:
    with open_resource(url, max_disk_bytes, stream_zip) as src:
        for table in src.tables:
            it = iter(table.rows)
            header = [h.strip().lstrip("﻿") for h in next(it, [])]
            for _ in zip(range(rows), it):
                pass
            return table.fmt, header
    raise NotTabular("no table", "no-tables")


# --- which resources may hold a table (from the metadata) -------------------------------

FORMATS = {"CSV": "csv", "TSV": "csv", "TXT": "csv", "ZIP": "zip", "CSV.ZIP": "zip", "CSV-ZIP": "zip",
           "XLSX": "xlsx", "XLS": "xls", "ODS": "ods", "PARQUET": "parquet", "JSON": "json", "XML": "xml"}
GEO_WORDS = ("shp", "shape", "kmz", "kml", "geojson", "gpkg", "wms", "wmts")
# Order of preference when the same table is published in several formats: the first is validated
# in full, the others only for their columns.
PRIORITY = {"csv": 0, "zip": 1, "parquet": 2, "xlsx": 3, "ods": 4, "xls": 5, "json": 6, "xml": 7}


def is_candidate(resource: dict) -> str | None:
    """The format a resource declares, when it may hold a table; None otherwise (maps, pages, PDFs)."""
    fmt = (resource.get("format") or "").upper().strip(". ")
    url = (resource.get("url") or "").lower().split("?")[0]
    if any(w in url.rsplit("/", 1)[-1] for w in GEO_WORDS) or fmt in ("GEOJSON", "SHP", "KMZ", "KML", "WMS", "SHP-ZIP"):
        return None
    if url.endswith(".zip"):
        return "zip"
    if fmt in FORMATS:
        return FORMATS[fmt]
    for ext, kind in ((".csv", "csv"), (".parquet", "parquet"), (".xlsx", "xlsx"), (".xls", "xls"), (".ods", "ods")):
        if url.endswith(ext):
            return kind
    return None


def safe_get(url: str) -> requests.Response:  # re-exported for tests that patch the network
    return ckan.get(url, stream=True)
