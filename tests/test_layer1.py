"""Layer 1 tests. No network: portals, files and the model are replaced by small fakes."""

import hashlib
import io
import json
import time
import zipfile
from datetime import datetime, timedelta, timezone

import pytest
import requests
import urllib3

from src import (census, ckan, config, dictionaries, drift, extract, linker, pdf_extract, report, safety, schemas,
                 tabular, types_map, validate, work)


# --- fakes -----------------------------------------------------------------------

class FakeRaw(io.BytesIO):
    decode_content = False


class FakeResponse:
    def __init__(self, data: bytes, headers=None):
        self.raw = FakeRaw(data)
        self.headers = headers or {"Content-Length": str(len(data)), "ETag": '"abc"'}

    def close(self):
        pass


REAL_PROBE = ckan.probe


@pytest.fixture(autouse=True)
def no_probes(monkeypatch):
    """The survey's delivery probes never reach the network in tests."""
    monkeypatch.setattr(ckan, "probe", lambda url: {"error": "ConnectionError"})
    monkeypatch.setattr(config, "VALIDATE_WORKERS", 1)     # validation in this process (fakes stay in place)


@pytest.fixture
def tmp_root(tmp_path, monkeypatch):
    """Every path the pipeline writes points into a temporary folder."""
    for name, rel in [("ROOT", ""), ("RESULTS", "results"), ("SCHEMAS", "schemas")]:
        monkeypatch.setattr(config, name, tmp_path / rel if rel else tmp_path)
    for name, rel in [("CENSUS_FILE", "census.json"), ("VALIDATION_FILE", "validation.json"),
                      ("EXTRACTION_FILE", "extraction.json"), ("QUEUE_FILE", "queue.json"),
                      ("DRIFT_FILE", "drift.json"), ("SUMMARY_FILE", "layer1_summary.json"),
                      ("RUN_LOG", "run_log.jsonl"), ("HISTORY_FILE", "history.json")]:
        monkeypatch.setattr(config, name, tmp_path / "results" / rel)
    monkeypatch.setattr(config, "DASHBOARD_FILE", tmp_path / "docs" / "data" / "layer1.json")
    return tmp_path


def serve(monkeypatch, files: dict[str, bytes]):
    """ckan.get / fetch_bytes answer from `files` (url -> content)."""
    monkeypatch.setattr(tabular.ckan, "get", lambda url, **kw: FakeResponse(files[url]))

    def fetch(url, limit):
        data = files[url]
        if len(data) > limit:
            raise ValueError("too large")
        return data, hashlib.sha256(data).hexdigest()

    monkeypatch.setattr(census.ckan, "fetch_bytes", fetch)
    monkeypatch.setattr(extract.ckan, "fetch_bytes", fetch)


IBAMA_DICT = ("nome_entidade;nome_atributo;tipificacao;datatype;descricao\n"
              "Termo;SEQ_TAD;público;INTEGER;Chave do termo\n"
              "Termo;DAT_TAD;público;DATA (DD/MM/AAAA);Data do termo\n"
              "Termo;UF;público;VARCHAR;Sigla\n").encode("cp1252")

RECIFE_DICT = json.dumps({"metadados": {
    "cabecalho": {"titulo": "Podas"},
    "recursos": [{"identificador": "http://dados.recife.pe.gov.br/dataset/podas/resource_edit/"
                                   "91357c64-ee2c-4139-9ce8-7f473c55f56d", "titulo": "podas_2022", "formato": "csv"}],
    "campos": [{"codigo": "cnpj", "descricao": "CNPJ", "tipo": "Char", "tamanho": 14, "valores_permitidos": ""},
               {"codigo": "valor", "descricao": "Valor", "tipo": "num", "tamanho": "", "valores_permitidos": ""}]}},
    ensure_ascii=False).encode("utf-8")


# --- types -----------------------------------------------------------------------

@pytest.mark.parametrize("declared,expected", [
    ("STRING", ("string", {})), ("TEXTO (STRING)", ("string", {})), ("Cadeia de caracteres", ("string", {})),
    ("char", ("string", {})), ("ENUM / VARCHAR", ("string", {})),
    ("INTEGER", ("integer", {})), ("NUMÉRICO (INTEIRO)", ("integer", {})), ("NUMBER(10)", ("integer", {})),
    ("NUMBER(10,2)", ("number", {})), ("num", ("number", {})), ("Numérico", ("number", {})),
    ("NÚMERO DECIMAL SEPARADOR PONTO", ("number", {"decimalChar": "."})),
    ("DATA (DD/MM/AA)", ("date", {"format": "%d/%m/%y"})), ("Data Simples", ("date", {})),
    ("DATA E HORA (DD/MM/AA HH:MM)", ("datetime", {"format": "%d/%m/%y %H:%M"})),
    ("TIMESTAMP", ("datetime", {})), ("DATATIME", ("datetime", {})), ("BOOLEAN", ("boolean", {})),
    ("json", ("object", {})), ("multipolygon", ("any", {})),
])
def test_declared_types_map_to_table_schema(declared, expected):
    ts_type, extra, _ = types_map.map_type(declared)
    assert (ts_type, extra) == expected


def test_unknown_or_missing_type_is_reported_not_guessed():
    assert types_map.map_type("")[2] == "type not declared"
    assert types_map.map_type("xyz")[0] == "any"
    assert "not recognised" in types_map.map_type("xyz")[2]


def test_field_descriptor_keeps_what_the_publisher_wrote():
    f = types_map.field_descriptor("UF", "VARCHAR", "2", "Sigla", "AC, AL")
    assert f["constraints"] == {"maxLength": 2} and f["declaredType"] == "VARCHAR"
    assert "Allowed values" in f["description"]
    assert "constraints" not in types_map.field_descriptor("V", "Numérico", "16,2")


# --- dictionaries ---------------------------------------------------------------

def test_reads_ibama_csv_dictionary_in_cp1252():
    d = dictionaries.read(IBAMA_DICT, "CSV")
    assert d.readable and d.reader == "csv"
    assert [f["name"] for f in d.parts[0].fields] == ["SEQ_TAD", "DAT_TAD", "UF"]
    assert d.parts[0].fields[1]["type"] == "DATA (DD/MM/AAAA)"


def test_reads_recife_json_dictionary_and_the_resources_it_names():
    d = dictionaries.read(RECIFE_DICT, "JSON")
    assert [f["name"] for f in d.parts[0].fields] == ["cnpj", "valor"]
    assert d.declared_resource_ids == ["91357c64-ee2c-4139-9ce8-7f473c55f56d"]
    assert d.declared_datasets == ["podas"]


def test_unreadable_dictionaries_get_a_reason():
    assert dictionaries.read(b"<!DOCTYPE html><html>portal error</html>", "CSV").error_kind == "html-page"
    assert dictionaries.read(b'{"metadados": {"campos": [', "JSON").error_kind == "malformed"
    assert dictionaries.read(b"a,b\n1,2\n", "CSV").error_kind == "no-field-table"


def test_reads_xlsx_dictionary_one_part_per_sheet():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Escolas"
    ws.append(["Dicionário de variáveis"])
    ws.append(["Nome da variável", "Descrição", "Tipo"])
    ws.append(["CO_ESCOLA", "Código", "Numérico"])
    ws.append(["NO_ESCOLA", "Nome", "Texto"])
    buf = io.BytesIO()
    wb.save(buf)
    d = dictionaries.read(buf.getvalue(), "XLSX")
    assert d.readable and d.parts[0].label == "Escolas"
    assert [f["name"] for f in d.parts[0].fields] == ["CO_ESCOLA", "NO_ESCOLA"]


def test_reads_xml_dictionary():
    xml = (b"<dicionario><campo><nome_atributo>A</nome_atributo><datatype>INTEGER</datatype></campo>"
           b"<campo><nome_atributo>B</nome_atributo><datatype>TEXTO</datatype></campo></dicionario>")
    d = dictionaries.read(xml, "XML")
    assert [f["name"] for f in d.parts[0].fields] == ["A", "B"]


def test_dictionary_in_the_resource_description_both_styles():
    a = "Intro.\n\nDicionário de dados:\n\n*\tSEQ_TAD – Chave.\n*\tNUM_TAD – Número.\n"
    b = "## Dicionário de Dados\n\n- **UF**: Sigla. Formato: texto;\n- **DATA_DE_EMISSAO**: Data. Formato: data no padrão DD/MM/AA;\n"
    assert [f["name"] for f in dictionaries.from_description(a)] == ["SEQ_TAD", "NUM_TAD"]
    fb = dictionaries.from_description(b)
    assert [(f["name"], f["type"]) for f in fb] == [("UF", "texto"), ("DATA_DE_EMISSAO", "data no padrão DD/MM/AA")]
    assert dictionaries.from_description("No dictionary here. * A – x") is None
    c = ("Dicionário de dados:\n* **Nº do Registro de Conversão:** Identificador. Formato: *Número*;\n"
         "* **Nome/Razão Social:** Nome. Formato: *Texto*;\n* Data Emissão - Data da emissão;\n")
    fc = dictionaries.from_description(c)
    assert [(f["name"], f["type"]) for f in fc] == [("Nº do Registro de Conversão", "Número"),
                                                   ("Nome/Razão Social", "Texto"), ("Data Emissão", "")]
    # labels written for people still match the technical column names after normalising
    m = validate.match_fields(["Nome/Razão Social", "Data Emissão"], ["NOME_RAZAO_SOCIAL", "DATA_EMISSAO"])
    assert m["missing"] == [] and m["spelling"] == {"Nome/Razão Social": "NOME_RAZAO_SOCIAL",
                                                    "Data Emissão": "DATA_EMISSAO"}


# --- PDF ------------------------------------------------------------------------

ANEEL_ROWS = [
    ["Nome do Campo", "Tipo do dado", "", "Tamanho\ndo Campo", "Descrição"],
    [None, None, None, "do Campo", None],
    ["DatGeracaoConjuntoDados", "Data Simples", "", None, "Data do processamento"],
    ["SigAgenteFiscalizador", "Cadeia de caracteres", "20", None, "Sigla do agente"],
    ["NumProcessoPunitivo", "Cadeia de caracteres", "50", "Número do processo. Pode ser"],   # page 2: shifted
    [None, None, None, "validado no SICNET."],                                               # continuation
    ["VlrPenalidade", "Numérico", "16,2", "Valor da penalidade."],
    ["DatLavramentoTE", "Data Simples", "", "Data de lavratura do termo"],
]


def test_pdf_rows_are_read_by_content_not_position():
    fields = pdf_extract.fields_from_table_rows(ANEEL_ROWS)
    assert [f["name"] for f in fields] == ["DatGeracaoConjuntoDados", "SigAgenteFiscalizador",
                                           "NumProcessoPunitivo", "VlrPenalidade", "DatLavramentoTE"]
    assert fields[1]["size"] == "20" and fields[3]["type"] == "Numérico" and fields[3]["size"] == "16,2"
    assert fields[2]["description"].endswith("validado no SICNET.")


def test_oracle_compares_extracted_names_with_the_header():
    fields = pdf_extract.fields_from_table_rows(ANEEL_ROWS)
    header = ["DatGeracaoConjuntoDados", "SigAgenteFiscalizador", "NumProcessoPunitivo", "VlrPenalidade",
              "DatLavraturaTE"]
    o = pdf_extract.oracle(fields, header)
    assert o["recall"] == 0.8 and o["precision"] == 0.8 and o["exact_match"] == 0.8
    assert o["missing_from_pdf"] == ["DatLavraturaTE"] and o["not_in_file"] == ["DatLavramentoTE"]
    assert o["similar_names"] == {}       # "lavramento" x "lavratura": another word (0.67), not a typo
    assert 0.8 < o["levenshtein"] < 1.0
    assert pdf_extract.agreement(o) == 0.8
    assert pdf_extract.oracle(fields, [])["available"] is False


def test_llm_output_is_bounded_data():
    text = json.dumps({"fields": [{"name": "A", "type": "int", "size": "", "description": "x" * 5000},
                                  {"name": "A", "type": "dup", "size": "", "description": ""},
                                  {"name": "", "type": "", "size": "", "description": ""}, "junk"]})
    out = pdf_extract.parse_llm(text)
    assert [f["name"] for f in out] == ["A"] and len(out[0]["description"]) == 600
    assert pdf_extract.parse_llm("not json") == []


def test_a_long_pdf_goes_to_the_model_whole_in_pieces(monkeypatch):
    lines = [f"CAMPO_{i:03d} texto descritivo do campo\n" for i in range(300)]
    text = "".join(lines)
    pieces = pdf_extract.chunks(text, 1000)
    assert "".join(pieces) == text and all(len(p) <= 1000 for p in pieces) and len(pieces) > 1
    assert pdf_extract.chunks("x" * 2500, 1000) == ["x" * 1000, "x" * 1000, "x" * 500]

    monkeypatch.setattr(pdf_extract, "pdf_text", lambda pdf: text)
    monkeypatch.setattr(config, "LLM_MAX_PDF_CHARS", 1000)

    class Client:
        def generate(self, system, prompt):
            names = [w for w in prompt.split() if w.startswith("CAMPO_")]
            return json.dumps({"fields": [{"name": n, "type": "", "size": "", "description": ""}
                                          for n in names + ["CAMPO_000"]]}), 1.0

    fields, meta = pdf_extract.llm(b"%PDF", Client())
    assert [f["name"] for f in fields] == [f"CAMPO_{i:03d}" for i in range(300)]   # none lost, none twice
    assert meta["pieces"] == len(pieces)

    class CutClient(Client):          # the answer is cut whenever the piece is larger than 600 characters
        def generate(self, system, prompt):
            self.last_done_reason = "length" if len(prompt) > 700 else "stop"
            return super().generate(system, prompt) if self.last_done_reason == "stop" else ('{"fields": [', 1.0)

    fields, meta = pdf_extract.llm(b"%PDF", CutClient())
    assert [f["name"] for f in fields] == [f"CAMPO_{i:03d}" for i in range(300)] and meta["pieces"] > len(pieces)


# --- linking --------------------------------------------------------------------

def _d(i, name, names=None, ids=None):
    return {"id": i, "name": name, "parts": [{"label": "", "names": names or []}], "declared_resource_ids": ids}


def test_links_ibama_dictionaries_by_name():
    dicts = [_d("d1", "Metadados - Termo de Suspensão"), _d("d2", "Metadados -  termo de suspensão - Enquadramento"),
             _d("d3", "Metadados - Termo de Suspensão - Enquadramento Complementar")]
    tables = [{"id": "r1", "name": "Termo de suspensão"}, {"id": "r2", "name": "Termo de Suspensão - Enquadramento"},
              {"id": "r3", "name": "Enquadramento Complementar"}]
    links = linker.link(dicts, tables, {})
    assert {k: v["dictionary"] for k, v in links.items()} == {"r1": "d1", "r2": "d2", "r3": "d3"}


def test_lone_generic_dictionary_describes_every_table_but_a_specific_one_does_not():
    tables = [{"id": "a", "name": "auto-infracao.csv"}, {"id": "b", "name": "auto-infracao-2020.csv"}]
    links = linker.link([_d("p", "Dicionário de dados")], tables, {})
    assert {v["method"] for v in links.values()} == {"single"} and set(links) == {"a", "b"}
    tables = [{"id": "a", "name": "Termos de embargo"}, {"id": "b", "name": "Coordenadas"}]
    assert set(linker.link([_d("p", "Metadados - Termos de embargo")], tables, {})) == {"a"}


def test_declared_id_and_header_beat_names():
    tables = [{"id": "r1", "name": "Podas 2022"}]
    dicts = [_d("byname", "Dicionário - Podas 2022"), _d("byid", "Dicionário - Erradicação", ids=["r1"])]
    assert linker.link(dicts, tables, {})["r1"]["method"] == "declared-id"
    dicts = [_d("byname", "Dicionário - Podas 2022"), _d("byheader", "Outro", names=["cnpj", "valor", "data"])]
    assert linker.link(dicts, tables, {"r1": ["cnpj", "valor", "data"]})["r1"]["dictionary"] == "byheader"
    assert linker.broken_declared_links([_d("x", "D", ids=["gone"])], {"r1"}) == {"x": ["gone"]}


# --- census ---------------------------------------------------------------------

def test_maturity_levels():
    assert census.level_of({}, None) == 0
    assert census.level_of({"described": 5}, None) == 1
    assert census.level_of({}, {"kind": "human"}) == 1
    assert census.level_of({}, {"kind": "machine", "readable": False}) == 1
    assert census.level_of({}, {"kind": "machine", "readable": True}) == 2
    assert census.level_of({"datastore": {"typed": False}}, {"kind": "machine", "readable": True}) == 2
    assert census.level_of({"datastore": {"typed": True}}, None) == 3
    assert census.level_of({"attached": True}, None) == 3


def test_datastore_with_only_text_fields_declares_nothing():
    assert schemas.from_datastore([{"id": "a", "type": "text"}, {"id": "b", "type": "text"}]) is None
    s = schemas.from_datastore([{"id": "a", "type": "text"}, {"id": "b", "type": "numeric"}])
    assert [f["type"] for f in s["fields"]] == ["string", "number"]


def _package(dict_url="https://p/d.csv", data_url="https://p/t.csv"):
    return {"name": "termo-de-doacao", "title": "Termo de doação", "organization": {"title": "IBAMA"},
            "resources": [
                {"id": "r-data", "name": "Termo de doação", "format": "CSV", "url": data_url},
                {"id": "r-dict", "name": "Metadados - Termo de doação", "format": "CSV", "url": dict_url},
                {"id": "r-json", "name": "Termo de doação", "format": "JSON", "url": "https://p/t.json"},
            ]}


def test_census_links_reads_and_writes_the_declared_schema(tmp_root, monkeypatch):
    serve(monkeypatch, {"https://p/d.csv": IBAMA_DICT})
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://p")
    result, events = census.run({}, packages=[_package()])
    ds = result["datasets"][0]
    assert [t["id"] for t in ds["tables"]] == ["r-data"]          # JSON copies are not validated
    t = ds["tables"][0]
    assert t["level"] == 2 and t["link"]["method"] == "name"
    assert ds["dictionaries"][0]["type_spellings"] == {"INTEGER": 1, "DATA (DD/MM/AAAA)": 1, "VARCHAR": 1}
    kind, schema = schemas.select("termo-de-doacao", "r-data")
    assert kind == "declared" and [f["name"] for f in schema["fields"]] == ["SEQ_TAD", "DAT_TAD", "UF"]
    assert schema["x5ltep"]["source"]["sha256"] == hashlib.sha256(IBAMA_DICT).hexdigest()
    assert events == []


def test_a_file_never_validated_is_linked_by_its_header_in_the_first_run(tmp_root, monkeypatch):
    # Before, the header was known only after a validation: a first run checked such a file against
    # the DataStore, and the next one against its dictionary, with the file unchanged.
    serve(monkeypatch, {"https://p/d.csv": IBAMA_DICT, "https://p/a.csv": b"SEQ_TAD;DAT_TAD;UF\n1;01/02/2022;PE\n",
                        "https://p/b.csv": b"X;Y\n1;2\n"})
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://p")
    pkg = {"name": "termo-de-doacao", "title": "T", "resources": [
        {"id": "r-a", "name": "Planilha A", "format": "CSV", "url": "https://p/a.csv"},
        {"id": "r-b", "name": "Outra coisa", "format": "CSV", "url": "https://p/b.csv"},
        {"id": "r-dict", "name": "Metadados - Termo de doação", "format": "CSV", "url": "https://p/d.csv"}]}
    result, _ = census.run({}, packages=[pkg])
    by_id = {t["id"]: t for t in result["datasets"][0]["tables"]}
    assert by_id["r-a"]["link"]["method"] == "header" and by_id["r-a"]["header_peeked"] == ["SEQ_TAD", "DAT_TAD", "UF"]
    assert by_id["r-b"]["link"] is None
    assert schemas.select("termo-de-doacao", "r-a")[0] == "declared"
    # a PDF dictionary has its oracle (the header) before the validation
    assert extract.header_for(["r-a"], {}, {"r-a": ["SEQ_TAD"]}) == (["SEQ_TAD"], True)
    assert extract.header_for(["r-a"], {}, {}) == (None, False)


def test_census_reports_declared_drift_and_removal(tmp_root, monkeypatch):
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://p")
    serve(monkeypatch, {"https://p/d.csv": IBAMA_DICT})
    census.run({}, packages=[_package()])
    changed = IBAMA_DICT + "Termo;MUNICIPIO;público;VARCHAR;Município\n".encode("cp1252")
    serve(monkeypatch, {"https://p/d.csv": changed})
    _, events = census.run({}, packages=[_package()])
    assert events[0]["kind"] == "declared" and events[0]["changes"]["fields_added"] == ["MUNICIPIO"]
    pkg = _package()
    pkg["resources"] = pkg["resources"][:1]          # the dictionary disappeared
    _, events = census.run({}, packages=[pkg])
    assert events[0]["changes"].get("schema_removed") is True
    assert schemas.select("termo-de-doacao", "r-data") == (None, None)


def test_portal_is_one_value(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PORTAL_FILE", tmp_path / "portal.json")
    (tmp_path / "portal.json").write_text('{"portal_url": "https://dados.recife.pe.gov.br/"}', encoding="utf-8")
    monkeypatch.delenv("CKAN_PORTAL_URL", raising=False)
    p = config.portal()
    assert p["portal_url"] == "https://dados.recife.pe.gov.br" and p["name"] == "dados.recife.pe.gov.br"
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://dadosabertos.aneel.gov.br")
    assert config.portal()["portal_url"] == "https://dadosabertos.aneel.gov.br"
    (tmp_path / "portal.json").write_text("{}", encoding="utf-8")
    monkeypatch.delenv("CKAN_PORTAL_URL")
    with pytest.raises(ValueError):
        config.portal()


# --- reading files as streams ------------------------------------------------------

def _zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


def test_reads_a_csv_stream_and_hashes_it(monkeypatch):
    data = "a;b\n1;ção\n2;x\n".encode("utf-8")
    serve(monkeypatch, {"u": data})
    with tabular.open_resource("u") as src:
        tables, rows = [], []
        for t in src.tables:          # rows must be read before the next table is asked for
            tables.append(t)
            rows = list(t.rows)
    assert tables[0].delimiter == ";" and tables[0].encoding == "utf-8" and rows[1] == ["1", "ção"]
    # the digest is complete once the generator is exhausted
    assert src.digest["sha256"] == hashlib.sha256(data).hexdigest() and src.digest["kind"] == "csv"
    assert src.digest["http"]["etag"] == '"abc"'


def test_detects_cp1252_and_counts_decode_errors(monkeypatch):
    serve(monkeypatch, {"u": "a,b\n1,ação\n".encode("cp1252")})
    with tabular.open_resource("u") as src:
        t = next(src.tables)
        list(t.rows)
    assert t.encoding == "cp1252" and t.delimiter == ","


def test_zip_members_are_tables(monkeypatch):
    data = _zip({"2020.csv": b"a;b\n1;2\n", "2021.csv": b"a;b;c\n1;2;3\n", "readme.pdf": b"%PDF"})
    serve(monkeypatch, {"u": data})
    with tabular.open_resource("u") as src:
        out = [validate.check_table(t, None) for t in src.tables]
    assert [o["member"] for o in out] == ["2020.csv", "2021.csv"]
    assert validate.summarise(out, None)["distinct_headers"] == 2
    assert src.digest["kind"] == "zip" and src.digest["members"] == 2


def test_html_or_zip_without_csv_is_not_tabular(monkeypatch):
    serve(monkeypatch, {"h": b"<!DOCTYPE html><html></html>", "z": _zip({"x.shp": b"0"})})
    for url in ("h", "z"):
        with pytest.raises(tabular.NotTabular):
            with tabular.open_resource(url) as src:
                list(src.tables)


def test_candidates_come_from_the_metadata():
    assert tabular.is_candidate({"format": "CSV", "url": "https://x/a.csv"}) == "csv"
    assert tabular.is_candidate({"format": "CSV", "url": "https://x/a_csv.zip"}) == "zip"
    assert tabular.is_candidate({"format": "ZIP", "url": "https://x/shapefile.zip"}) is None
    assert tabular.is_candidate({"format": "PDF", "url": "https://x/a.pdf"}) is None


# --- validation --------------------------------------------------------------------

class Table:
    def __init__(self, text: str, member=None):
        import csv as _csv
        self.member, self.encoding, self.delimiter = member, "utf-8", ";"
        self.rows = iter(list(_csv.reader(io.StringIO(text), delimiter=";")))
        self.decode_errors = [0]


DECLARED = {"fields": [types_map.field_descriptor("ID", "INTEGER"),
                       types_map.field_descriptor("DATA", "Data Simples"),
                       types_map.field_descriptor("VALOR", "Numérico"),
                       types_map.field_descriptor("UF", "VARCHAR", "2"),
                       types_map.field_descriptor("Descrição", "TEXTO"),
                       types_map.field_descriptor("SUMIU", "TEXTO")]}
CSV_TEXT = ("ID;DATA;VALOR;UF;DESCRICAO;EXTRA\n"
            "1;03/05/2022;1.234,50;PE;Maria da Silva;x\n"
            "2;04/05/2022;,99;PE;João;x\n"
            "3;31/02/2022;10,00;PER;José;x\n"
            "abc;05/05/2022;7,5;SP;;x\n")


def test_conformance_counts_errors_and_infers_undeclared_formats():
    out = validate.check_table(Table(CSV_TEXT), DECLARED)
    c = out["conformance"]
    assert c["matches"]["missing"] == ["SUMIU"] and c["matches"]["undeclared"] == ["EXTRA"]
    assert c["matches"]["spelling"] == {"Descrição": "DESCRICAO"}
    assert c["matches"]["similar_names"] == {}                     # SUMIU and EXTRA are not alike
    assert validate.match_fields(["IdeNuceloCEG", "UF"], ["IdeNucleoCEG", "UF", "Outra"])["similar_names"] ==         {"IdeNuceloCEG": "IdeNucleoCEG"}
    assert c["inferred_formats"]["DATA"] == {"format": "%d/%m/%Y"}
    assert c["inferred_formats"]["VALOR"] == {"decimalChar": ",", "groupChar": "."}
    errs = {f: {k: e["count"] for k, e in kinds.items()} for f, kinds in c["errors"].items()}
    assert errs == {"ID": {"type": 1}, "DATA": {"type": 1}, "UF": {"maxLength": 1}}
    assert c["errors"]["ID"]["type"]["rows"] == [4]
    s = validate.summarise([out], DECLARED)["conformance"]
    assert s["pass"] is False and s["cells_with_errors"] == 3


def test_no_cell_value_is_ever_kept():
    out = validate.summarise([validate.check_table(Table(CSV_TEXT), DECLARED)], DECLARED)
    dumped = json.dumps(out, ensure_ascii=False)
    for value in ("Maria da Silva", "João", "1.234,50", "31/02/2022", "PER", "abc"):
        assert value not in dumped


def test_a_parse_error_keeps_no_piece_of_the_file(tmp_root, monkeypatch):
    # A JSON exported with NaN is not JSON: the parser's message quotes the text around the fault.
    data = b'[{"NOME": "Maria da Silva", "DAP": NaN}]'
    serve(monkeypatch, {"https://p/t.json": data})
    t = {"id": "r1", "name": "T", "url": "https://p/t.json", "format": "JSON"}
    entry, _, _ = work.validate_one("ds", t, None)
    assert entry["status"] == "error"
    assert "Maria" not in json.dumps(entry, ensure_ascii=False)
    # a record written before first_seen existed is not "new" when it is validated again
    serve(monkeypatch, {"https://p/t.csv": b"A;B\n1;2\n"})
    old = {"status": "ok", "checked_at": "2026-09-01T00:00:00+00:00"}
    again, _, _ = work.validate_one("ds", {**t, "url": "https://p/t.csv", "format": "CSV"}, old)
    assert again["first_seen"] == "2026-09-01T00:00:00+00:00"
    import requests
    assert "https://p/x" in safety.error_text(requests.HTTPError("404 Client Error: Not Found for url: https://p/x"))


def test_a_clean_file_passes():
    text = "ID;DATA;VALOR;UF;Descrição;SUMIU\n1;2022-05-03;1.5;PE;a;b\n2;2022-05-04;2;SP;c;d\n"
    s = validate.summarise([validate.check_table(Table(text), DECLARED)], DECLARED)["conformance"]
    assert s["pass"] is True and s["error_rate"] == 0


def test_observed_types():
    assert validate.infer_type(["1", "2", ""])["type"] == "integer"
    assert validate.infer_type(["01", "2"])["type"] == "string"          # leading zeros: a code
    assert validate.infer_type(["1,5", ",99"]) == {"type": "number", "decimalChar": ","}
    assert validate.infer_type(["2022-01-03"]) == {"type": "date", "format": "%Y-%m-%d"}
    assert validate.infer_type(["S", "N"])["type"] == "boolean"
    assert validate.infer_type([""]) == {"type": "any"}


# --- work queue -----------------------------------------------------------------------

def _census_one(url="https://p/t.csv", lm="2026-01-01"):
    return {"datasets": [{"name": "ds", "dictionaries": [], "tables": [
        {"id": "r1", "name": "T", "url": url, "last_modified": lm, "size": 10, "level": 0}]}]}


def test_queue_reasons(tmp_root):
    now = datetime(2026, 10, 1, tzinfo=timezone.utc)
    recent = (now - timedelta(days=3)).isoformat()
    prev = {"r1": {"ckan": {"url": "https://p/t.csv", "last_modified": "2026-01-01", "size": 10},
                   "schema_fingerprint": None, "validated_at": recent, "status": "ok",
                   "http": {"etag": '"a"'}}}
    no_change = lambda url: {"etag": '"a"'}
    assert work.plan(_census_one(), {}, now)[0]["reason"] == "new"
    assert work.plan(_census_one(), prev, now, head=no_change) == []
    assert work.plan(_census_one(lm="2026-02-01"), prev, now)[0]["reason"] == "changed"
    assert work.plan(_census_one(), prev, now, head=lambda url: {"etag": '"b"'})[0]["reason"] == "changed"
    old = {"r1": {**prev["r1"], "validated_at": (now - timedelta(days=40)).isoformat()}}
    assert work.plan(_census_one(), old, now, head=no_change)[0]["reason"] == "rotation"
    before_kinds = {"r1": {**prev["r1"], "status": "not-tabular", "error": "zip without CSV members"}}
    assert work.plan(_census_one(), before_kinds, now, head=no_change)[0]["reason"] == "reader"
    kind_known = {"r1": {**before_kinds["r1"], "not_tabular_kind": "no-tables"}}
    assert work.plan(_census_one(), kind_known, now, head=no_change) == []
    census = _census_one()
    census["datasets"][0]["tables"][0]["distributions"] = [{"id": "r2", "format": "XLSX", "url": "https://p/t.xlsx"}]
    assert work.plan(census, prev, now, head=no_change)[0]["reason"] == "reader"
    compared = {"r1": {**prev["r1"], "distributions": []}}
    assert work.plan(census, compared, now, head=no_change) == []
    schemas.write(schemas.path("ds", "r1", "declared"), DECLARED)
    assert work.plan(_census_one(), prev, now, head=no_change)[0]["reason"] == "schema"


def test_a_file_killed_mid_validation_is_not_retried_in_the_same_chain(tmp_root, monkeypatch):
    saved = []

    def killed(*args):
        raise KeyboardInterrupt      # stands for the job being killed by its time limit

    monkeypatch.setattr(work, "validate_one", killed)
    serve(monkeypatch, {"https://p/t.csv": b"ID;UF\n1;PE\n"})
    validation = {}
    with pytest.raises(KeyboardInterrupt):
        work.run_batch(_census_one(), [{"id": "r1", "dataset": "ds", "reason": "new"}], validation, 5,
                       on_progress=lambda out: saved.append((list(out["done"]), dict(validation))))
    done, state = saved[-1]
    assert done == ["r1"] and state["r1"]["status"] == "error"
    assert "time limit" in state["r1"]["error"]


def test_several_files_are_validated_at_once_and_a_large_spreadsheet_alone(tmp_root, monkeypatch):
    # 03/10/2026: the validation of a zip is CPU-bound (downloaded at 6-23 MB/s, validated at 1.4-1.9
    # MB/s), on a runner with 4 vCPUs.
    import tempfile
    import threading
    from concurrent.futures import ThreadPoolExecutor

    lock, now, peak, order = threading.Lock(), [0], [0], []
    monkeypatch.setattr(work, "_fetch", lambda url, limit: {"dir": tempfile.mkdtemp(), "digest": {}})

    def fake(dataset, t, prev, local=None):
        with lock:
            now[0] += 1
            peak[0] = max(peak[0], now[0])
            order.append((t["id"], now[0]))
        time.sleep(0.05)
        with lock:
            now[0] -= 1
        return {"dataset": dataset, "status": "ok"}, None, []

    monkeypatch.setattr(work, "validate_one", fake)
    big = 2 * config.MAX_ZIP_BYTES
    census_ = {"datasets": [{"name": "ds", "dictionaries": [], "tables": [
        {"id": f"r{i}", "name": "T", "url": "u", "candidate": "csv", "size": 10} for i in range(6)]
        + [{"id": "x", "name": "X", "url": "u", "candidate": "xlsx", "size": big}]}]}
    queue = [{"id": f"r{i}", "dataset": "ds", "reason": "new"} for i in range(3)] + [{"id": "x", "dataset": "ds", "reason": "new"}] \
        + [{"id": f"r{i}", "dataset": "ds", "reason": "new"} for i in range(3, 6)] + [{"id": "gone", "dataset": "ds", "reason": "new"}]
    validation = {}
    out = work.run_batch(census_, queue, validation, 5, workers=3, executor=ThreadPoolExecutor(3))
    assert sorted(out["done"]) == sorted(["r0", "r1", "r2", "r3", "r4", "r5", "x", "gone"])
    assert peak[0] == 3 and dict(order)["x"] == 1           # the spreadsheet that may not fit ran alone
    assert all(validation[f"r{i}"]["status"] == "ok" for i in range(6))
    # past the budget nothing new starts
    assert work.run_batch(census_, queue, {}, -1, workers=3, executor=ThreadPoolExecutor(3))["done"] == []


def test_one_connection_per_server_and_the_validation_from_disk(tmp_root, monkeypatch):
    # 03/10/2026: three connections at once to ANEEL's server made each one 5 to 10 times slower
    import threading
    from concurrent.futures import ThreadPoolExecutor

    lock, open_, peak = threading.Lock(), {}, {}
    body = b"ID;UF\n1;PE\n"

    def fetch(url, folder, limit):
        host = url.split("/")[2]
        with lock:
            open_[host] = open_.get(host, 0) + 1
            peak[host] = max(peak.get(host, 0), open_[host])
            peak["all"] = max(peak.get("all", 0), sum(open_.values()))
        time.sleep(0.05)
        (folder / "resource").write_bytes(body)
        with lock:
            open_[host] -= 1
        return {"timing": {"download_s": 0.05, "network_s": 0.05}, "http": {}, "bytes": len(body),
                "sha256": hashlib.sha256(body).hexdigest()}

    monkeypatch.setattr(tabular, "fetch", fetch)
    urls = [f"https://a/{i}.csv" for i in range(4)] + [f"https://b/{i}.csv" for i in range(2)]
    census_ = {"datasets": [{"name": "ds", "dictionaries": [], "tables": [
        {"id": f"r{i}", "name": "T", "url": u, "candidate": "csv", "size": 10} for i, u in enumerate(urls)]}]}
    validation = {}
    out = work.run_batch(census_, [{"id": f"r{i}", "dataset": "ds", "reason": "new"} for i in range(6)],
                         validation, 5, workers=3, executor=ThreadPoolExecutor(3))
    assert sorted(out["done"]) == [f"r{i}" for i in range(6)]
    assert peak["a"] == 1 and peak["b"] == 1 and peak["all"] == 2
    assert all(validation[f"r{i}"]["status"] == "ok" and validation[f"r{i}"]["header"] == ["ID", "UF"] for i in range(6))
    assert validation["r0"]["seconds"] >= 0.05 and validation["r0"]["timing"]["network_s"] == 0.05


def test_a_file_larger_than_its_share_of_the_disk_is_read_over_the_network(tmp_root, monkeypatch):
    body = b"ID;UF\n" + b"1;PE\n" * 20
    serve(monkeypatch, {"https://p/t.csv": body})
    monkeypatch.setattr(config, "MAX_ZIP_BYTES", 2 * len(body) - 2)      # share: just under the file
    validation = {}
    census_ = _census_one()
    census_["datasets"][0]["tables"][0]["size"] = None                    # the portal does not say
    out = work.run_batch(census_, [{"id": "r1", "dataset": "ds", "reason": "new"}], validation, 5)
    assert out["done"] == ["r1"] and validation["r1"]["status"] == "ok"
    assert validation["r1"]["sha256"] == hashlib.sha256(body).hexdigest()
    assert "download_s" not in validation["r1"]["timing"]                   # read while it streamed


def test_a_zip_is_validated_while_it_downloads_in_the_order_read_from_disk(tmp_root, monkeypatch):
    # stored as b.csv, a.csv: from disk the members are read by name, and the first header is the one
    # compared for drift, so the streamed reading gives them in the same order
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("b.csv", b"B1;B2\n1;2\n")
        zf.writestr("a.csv", b"A1;A2\n1;2\n")
    serve(monkeypatch, {"https://p/t.zip": buf.getvalue()})
    entry, observed, _ = work.validate_one("ds", {"id": "r1", "name": "T", "url": "https://p/t.zip"}, None)
    assert entry["header"] == ["A1", "A2"] and entry["sha256"] == hashlib.sha256(buf.getvalue()).hexdigest()
    assert entry["timing"]["network_s"] >= 0 and "download_s" not in entry["timing"]
    assert tabular.member_order("n.zip/m.zip/b.csv") < tabular.member_order("n.zip-old.csv")


def test_the_parallel_validation_runs_in_separate_processes(tmp_root):
    # the real path: spawned processes, a file served over HTTP on this machine
    import http.server
    import threading

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            body = b"ID;UF\n1;PE\n"
            self.send_response(200)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *a):
            pass

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        url = f"http://127.0.0.1:{server.server_port}/t.csv"
        census_ = {"datasets": [{"name": "ds-test-parallel", "dictionaries": [], "tables": [
            {"id": f"p{i}", "name": "T", "url": url, "candidate": "csv", "size": 10} for i in range(3)]}]}
        validation = {}
        out = work.run_batch(census_, [{"id": f"p{i}", "dataset": "ds-test-parallel", "reason": "new"} for i in range(3)],
                             validation, 5, workers=2)
    finally:
        server.shutdown()
    assert sorted(out["done"]) == ["p0", "p1", "p2"]
    assert all(validation[f"p{i}"]["status"] == "ok" and validation[f"p{i}"]["header"] == ["ID", "UF"] for i in range(3))


def test_batch_validates_and_records_observed_schema(tmp_root, monkeypatch):
    serve(monkeypatch, {"https://p/t.csv": b"ID;UF\n1;PE\n"})
    validation = {}
    out = work.run_batch(_census_one(), [{"id": "r1", "dataset": "ds", "reason": "new"}], validation, 5)
    assert out["done"] == ["r1"] and validation["r1"]["status"] == "ok"
    assert validation["r1"]["header"] == ["ID", "UF"] and validation["r1"]["sha256"]
    assert out["observed"][("ds", "r1")]["x5ltep"]["delimiter"] == ";"


# --- drift --------------------------------------------------------------------------

def _obs(names, delimiter=";"):
    return {"fields": [{"name": n} for n in names], "x5ltep": {"delimiter": delimiter, "encoding": "utf-8",
                                                               "distinct_headers": 1, "observed_at": "t"}}


def test_observed_drift():
    t = {"id": "r1", "name": "T"}
    assert drift.compare("ds", t, None, _obs(["a"])) == []                         # baseline
    assert drift.compare("ds", t, _obs(["a", "b"]), _obs(["a", "b"])) == []
    e = drift.compare("ds", t, _obs(["a", "b"]), _obs(["a", "c"]))[0]
    assert e["changes"] == {"columns_added": ["c"], "columns_removed": ["b"]}
    assert drift.compare("ds", t, _obs(["a", "b"]), _obs(["b", "a"]))[0]["changes"] == {"columns_reordered": True}
    e = drift.compare("ds", t, _obs(["a"]), _obs(["a"], ","))[0]
    assert e["changes"] == {"delimiter": {"before": ";", "after": ","}}


def test_drift_log_is_append_only_and_idempotent(tmp_root):
    e = drift.compare("ds", {"id": "r1"}, _obs(["a"]), _obs(["b"]))
    assert drift.append(e) == 1 and drift.append(e) == 0 and len(drift.load()) == 1


def test_issue_text_is_neutralised():
    e = {"id": "x", "kind": "observed", "dataset": "ds", "resource_id": "r1",
         "resource_name": "<img src=x> @someone #12 [a](https://evil)",
         "changes": {"columns_added": ["`rm -rf`", "| table |"]}}
    body = drift.issue_body(e, "https://p", "https://d/")
    title = drift.issue_title(e)
    for bad in ("<img", "@someone", "#12", "](https://evil)"):
        assert bad not in body + title
    assert body.startswith("<!-- l1-drift:x -->")


# --- PDF extraction stages -------------------------------------------------------------

def _pdf_census(header_known=True):
    c = {"datasets": [{"name": "ds", "tables": [{"id": "r1", "name": "auto.csv"}], "dictionaries": [
        {"id": "pdf1", "name": "Dicionário de dados", "format": "PDF", "kind": "human", "url": "https://p/d.pdf",
         "sha256": hashlib.sha256(b"%PDF-1").hexdigest(), "linked_resources": ["r1"]}]}]}
    validation = {"r1": {"status": "ok", "header": ["A", "B", "C"]}} if header_known else {}
    return c, validation


def test_pdf_waits_for_the_oracle_then_extracts(tmp_root, monkeypatch):
    serve(monkeypatch, {"https://p/d.pdf": b"%PDF-1"})
    c, _ = _pdf_census(header_known=False)
    assert extract.tasks(c, {}, {}, "deterministic") == []
    c, validation = _pdf_census()
    todo = extract.tasks(c, validation, {}, "deterministic")
    monkeypatch.setattr(extract.pdf_extract, "deterministic",
                        lambda pdf: [{"name": n, "type": "TEXTO", "size": "", "description": "", "allowed": ""}
                                     for n in ("A", "B", "C")])
    rec = {}
    out = extract.run_deterministic(todo, rec)
    assert rec["pdf1"]["outcome"] == "extracted" and ("ds", "r1") in out["extracted"]
    assert out["extracted"][("ds", "r1")]["x5ltep"]["status"] == "extracted"
    assert extract.tasks(c, validation, rec, "deterministic") == []                # same PDF: done


def test_partial_agreement_goes_to_the_llm_and_then_to_people(tmp_root, monkeypatch):
    serve(monkeypatch, {"https://p/d.pdf": b"%PDF-1"})
    c, validation = _pdf_census()
    monkeypatch.setattr(extract.pdf_extract, "deterministic",
                        lambda pdf: [{"name": "A", "type": "", "size": "", "description": "", "allowed": ""}])
    rec = {}
    extract.run_deterministic(extract.tasks(c, validation, {}, "deterministic"), rec)
    assert rec["pdf1"]["outcome"] == "llm-needed"

    class Client:
        def info(self):
            return {"model_digest": "sha256:x"}

    monkeypatch.setattr(extract.pdf_extract, "llm", lambda pdf, client: (
        [{"name": n, "type": "", "size": "", "description": "", "allowed": ""} for n in ("A", "B", "C")], {"seconds": 1}))
    todo = extract.tasks(c, validation, rec, "llm")
    out = extract.run_llm(todo, rec, Client(), 5)
    assert rec["pdf1"]["outcome"] == "suggested" and rec["pdf1"]["stage"] == "llm"
    s = out["suggested"][("ds", "r1")]
    assert s["x5ltep"]["status"] == "suggested" and s["x5ltep"]["source"]["model"] == config.LLM_MODEL
    assert extract.tasks(c, validation, rec, "llm") == []                           # not asked twice


def test_suggested_schemas_are_never_used_by_the_validation(tmp_root):
    s = schemas.from_fields([{"name": "A"}], "suggested", {})
    schemas.write(schemas.path("ds", "r1", "extracted"), s)
    assert schemas.select("ds", "r1") == (None, None)
    s["x5ltep"]["status"] = "verified"
    schemas.write(schemas.path("ds", "r1", "extracted"), s)
    assert schemas.select("ds", "r1")[0] == "extracted"


# --- safety ---------------------------------------------------------------------------

def _artifact(tmp_path, files: dict[str, object]):
    src = tmp_path / "artifact"
    for rel, content in files.items():
        p = src / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(content if isinstance(content, str) else json.dumps(content), encoding="utf-8")
    return src


def test_accepts_the_expected_files(tmp_root):
    c, _ = _pdf_census()
    obs = {**_obs(["A"]), "x5ltep": {"status": "observed"}}
    src = _artifact(tmp_root, {"results/validation.json": {"r1": {"status": "ok"}}, "results/done.json": ["r1"],
                               "schemas/ds/r1.observed.json": obs, "chain.json": {},
                               "suggestions/ds/r1.extracted.json": schemas.from_fields([{"name": "A"}], "suggested", {})})
    accepted = safety.accept_artifact(src, tmp_root / "repo", c, [])
    assert "schemas/ds/r1.observed.json" in accepted and (tmp_root / "repo/suggestions/ds/r1.extracted.json").exists()


def test_schemas_of_unknown_resources_are_never_written(tmp_root):
    c, _ = _pdf_census()
    src = _artifact(tmp_root, {"schemas/other/r1.observed.json": {"fields": [], "x5ltep": {"status": "observed"}}})
    assert safety.accept_artifact(src, tmp_root / "repo", c, []) == []
    assert not (tmp_root / "repo/schemas/other").exists()


@pytest.mark.parametrize("files", [
    {".github/workflows/x.yml": "on: push"},
    {"schemas/Bad Name/r1.observed.json": {"fields": [], "x5ltep": {"status": "observed"}}},
    {"suggestions/ds/r1.observed.json": {"fields": [{"name": "A"}], "x5ltep": {"status": "observed"}}},
    {"results/validation.json": {"r1": {"status": "pwned"}}},
    {"results/drift.json": []},
])
def test_refuses_anything_else(tmp_root, files):
    c, _ = _pdf_census()
    committed = [{"id": "e1", "kind": "declared", "resource_id": "r1"}]
    with pytest.raises(ValueError):
        safety.accept_artifact(_artifact(tmp_root, files), tmp_root / "repo", c, committed)


# --- report --------------------------------------------------------------------------

def test_summary_levels_rate_and_findings(tmp_root):
    c = {"portal": {"portal_url": "https://p", "name": "P", "title": "P"}, "generated_at": "t", "datasets": [
        {"name": "a", "title": "A", "url": "u", "dictionaries": [
            {"id": "d1", "name": "Metadados", "format": "CSV", "kind": "machine", "readable": True,
             "type_spellings": {"TEXTO": 2, "xyz": 1}, "unmapped_types": ["xyz"], "date_fields": 1,
             "date_fields_with_format": 0, "linked_resources": ["t1"]},
            {"id": "d2", "name": "Metadados 2", "format": "CSV", "kind": "machine", "readable": False,
             "error_kind": "html-page", "linked_resources": []}],
         "tables": [{"id": "t1", "name": "T1", "url": "u", "level": 3, "link": {"method": "name", "dictionary": "d1"}},
                    {"id": "t2", "name": "T2", "url": "u", "level": 0, "link": None}]}]}
    conf = {"pass": True, "missing": [], "undeclared": [], "spelling": {}, "inferred_formats": {"D": {"format": "%d/%m/%Y"}},
            "error_rate": 0, "errors_by_field": {}}
    v = {"t1": {"status": "ok", "schema_kind": "declared", "summary": {"conformance": conf, "encodings": ["utf-8"],
                                                                        "delimiters": [";"], "rows": 2}},
         "t2": {"status": "ok", "summary": {"encodings": ["cp1252"], "delimiters": [","], "rows": 2}}}
    s = report.build_summary(c, v, {}, [])
    assert s["tables"]["by_level"] == {"0": 1, "1": 0, "2": 0, "3": 0, "4": 1}
    assert s["datasets"]["by_level"]["0"] == 1                           # weakest table
    assert s["l1_rate"] == 1.0 and s["l1_pass"] is True
    assert s["conformance_by_source"] == {"declared": {"verifiable": 1, "conformant": 1, "l1_rate": 1.0}}
    f = s["findings"]
    assert f["dictionaries"]["unreadable_by_reason"] == {"html-page": 1} and f["dictionaries"]["orphans"] == 1
    assert f["types"]["distinct_spellings"] == 2 and f["types"]["fields_with_recognised_type"] == round(2 / 3, 4)
    assert f["files_vs_dictionaries"]["formats_inferred_from_data"] == {"format": 1}
    d = report.dashboard_data(c, v, {}, s)
    assert d["datasets"][0]["tables"][0]["level"] == 4


def test_file_delivery_is_measured_per_server_and_described_against_http(tmp_root, monkeypatch):
    # 03/10/2026: IBAMA (Azure Blob) and ANEEL (CKAN, nginx) answer a byte-range request with 206;
    # Recife's download service answers with the whole file, without ETag or Last-Modified, and
    # serves CSV as application/octet-stream.
    class Resp:
        def __init__(self, status, headers, url):
            self.status_code, self.headers, self.url, self.closed = status, headers, url, False

        def close(self):
            self.closed = True

    sent = []

    def get(url, headers=None, **kw):
        sent.append(headers.get("Range"))
        return Resp(200, {"Content-Type": "application/octet-stream", "Content-Length": "9", "Server": "istio-envoy"},
                    "https://storage.r/x.csv")

    monkeypatch.setattr(ckan.SESSION, "get", get)
    p = REAL_PROBE("https://r/x.csv")
    assert sent == ["bytes=0-99"] and p["status"] == 200 and p["served_by"] == "storage.r"
    assert p["content_type"] == "application/octet-stream" and not p["etag"] and not p["last_modified"]
    # the survey probes a few files per server, from different datasets
    monkeypatch.setattr(ckan, "probe", lambda url: {"status": 200 if "//r/" in url else 206, "served_by": "storage.r",
                                                    "server": "istio-envoy", "accept_ranges": None, "etag": False,
                                                    "last_modified": False, "content_length": True,
                                                    "content_type": "application/octet-stream"})
    ds = lambda i, host: {"name": f"d{i}", "tables": [{"id": f"t{i}", "url": f"https://{host}/{i}.csv", "candidate": "csv", "level": 0},
                                                     {"id": f"u{i}", "url": f"https://{host}/u{i}.csv", "candidate": "csv", "level": 0}]}
    probes = census.probe_servers([ds(i, "r") for i in range(5)] + [ds(9, "b")])
    assert [len(probes["by_host"][h]) for h in ("b", "r")] == [1, census.config.DELIVERY_PROBES]
    # the report sets what was observed beside what HTTP provides, per server
    c = {"portal": {"portal_url": "https://r", "name": "R", "title": "R"}, "generated_at": "t", "delivery": probes,
         "datasets": [{**ds(i, "r"), "title": "", "url": "", "dictionaries": []} for i in range(2)]}
    http = {"etag": None, "last_modified": None, "content_length": 9, "content_type": "application/octet-stream",
            "served_by": "storage.r"}
    v = {"t0": {"status": "ok", "host": "r", "kind": "csv", "http": http, "summary": {"rows": 1}},
         "u0": {"status": "ok", "host": "r", "kind": "csv", "http": {**http, "content_type": "text/csv", "etag": '"e"'},
                "summary": {"rows": 1}},
         "t1": {"status": "error", "host": "r", "error": "404 Client Error: Not Found for url"},
         "u1": {"status": "error", "host": "r", "error": "TooManyRedirects: Exceeded 30 redirects."}}
    d = report.build_summary(c, v, {}, [])["findings"]["delivery"]
    r = d["hosts"]["r"]
    assert r["range"] == "ignored" and r["served_by"] == ["storage.r"] and r["server"] == ["istio-envoy"]
    assert {s["id"]: (s["n"], s["of"]) for s in d["suggestions"] if s["host"] == "r"} == {
        "broken-links": (1, 4), "redirect-loops": (1, 4), "byte-ranges": (3, 3), "validators": (1, 2), "media-type": (1, 2)}
    assert d["hosts"]["b"]["range"] == "honoured" and not [s for s in d["suggestions"] if s["host"] == "b" and s["id"] == "byte-ranges"]


def test_status_badge(tmp_root):
    s = {"tables": {"queue_remaining": 7}, "conformant": 3, "verifiable": 9}
    assert report.status_from_summary(s, True) == ("validating", {"left": 7})
    assert report.status_from_summary(s, False) == ("done", {"conformant": 3, "verifiable": 9})
    report.write_status("done", conformant=3, verifiable=9)
    pt = json.loads((tmp_root / "docs/data/status.pt.json").read_text(encoding="utf-8"))
    assert pt["label"] == "Camada 1" and "3/9" in pt["message"]


# --- documentation -----------------------------------------------------------------------

def _doc_structure(text: str) -> dict:
    headings, fences, in_code = [], 0, False
    for line in text.splitlines():
        if line.startswith("```"):
            fences += 1
            in_code = not in_code
        elif not in_code and line.startswith("#"):
            headings.append(len(line) - len(line.lstrip("#")))
    return {"headings": headings, "code_blocks": fences // 2}


def test_readme_and_leiame_stay_parallel():
    root = config.Path(__file__).resolve().parent.parent
    readme = (root / "README.md").read_text(encoding="utf-8")
    leiame = (root / "LEIAME.md").read_text(encoding="utf-8")
    assert _doc_structure(readme) == _doc_structure(leiame)
    assert "(LEIAME.md)" in readme and "(README.md)" in leiame


# --- model selection ---------------------------------------------------------------------

class FakeGet:
    def __init__(self, payload=None, error=None):
        self.payload, self.error = payload, error

    def __call__(self, url, timeout=None):
        if self.error:
            raise self.error
        return type("R", (), {"json": lambda _self: self.payload})()


def test_model_follows_the_benchmark_or_falls_back(monkeypatch):
    from src import model_select

    monkeypatch.setattr(config, "LLM_MODEL", "auto")
    monkeypatch.setattr(config, "LLM_THINK", "")
    rec = {"use": {"backend": "ollama", "model": "qwen3:4b", "digest": "a" * 64, "options": {"think": False}}}
    monkeypatch.setattr(model_select.requests, "get", FakeGet(rec))
    assert model_select.resolve() == {"model": "qwen3:4b", "think": "false", "source": "benchmark",
                                      "digest": "a" * 64, "benchmark_generated_at": None}
    monkeypatch.setattr(config, "LLM_MODEL", "auto")
    monkeypatch.setattr(model_select.requests, "get", FakeGet({"use": {"model": "x; rm -rf /"}}))
    assert model_select.resolve()["source"] == "fallback"
    monkeypatch.setattr(config, "LLM_MODEL", "auto")
    monkeypatch.setattr(model_select.requests, "get", FakeGet(error=model_select.requests.ConnectionError()))
    assert model_select.resolve()["model"] == config.FALLBACK_MODEL
    monkeypatch.setattr(config, "LLM_MODEL", "gemma3:4b")
    assert model_select.resolve()["source"] == "pinned"
    monkeypatch.setattr(config, "LLM_MODEL", "bad name!")
    with pytest.raises(ValueError):
        model_select.resolve()


def test_coverage_funnel_and_reasons(tmp_root):
    conf = lambda **kw: {"pass": False, "missing": [], "undeclared": [], "error_rate": 0.0, **kw}
    tables = [
        ({"url": "https://a.gov/x.csv"}, {"status": "ok", "summary": {"conformance": {**conf(), "pass": True}}}),
        ({"url": "https://a.gov/y.csv"}, {"status": "ok", "summary": {"conformance": conf(missing=["X"])}}),
        ({"url": "https://a.gov/z.csv"}, {"status": "ok", "summary": {"conformance": conf(error_rate=0.5)}}),
        ({"url": "https://a.gov/w.csv"}, {"status": "ok", "summary": {}}),
        ({"url": "https://geo.gov/ows"}, {"status": "error", "error": "HTTPError: 404 Client Error: Not Found"}),
        ({"url": ""}, {"status": "error", "error": "MissingSchema: Invalid URL ''"}),
        ({"url": "https://a.gov/e.csv"}, {"status": "empty"}),
        ({"url": "https://a.gov/p.csv"}, None),
    ]
    c = report.coverage(tables)
    assert (c["files"], c["read"], c["checked"], c["conform"], c["read_without_schema"]) == (8, 4, 3, 1, 1)
    assert (c["not_downloaded"], c["empty"], c["pending"]) == (2, 1, 1)
    assert c["not_downloaded_by_host"] == {"HTTP 404 · geo.gov": 1, "no URL · —": 1}
    assert (c["fail_missing_fields"], c["fail_cells_over_limit"], c["fail_only_cells"]) == (1, 1, 1)
    b = report.drift_baseline({"r1": {"status": "ok", "header": ["a"], "validated_at": "2026-10-01T10:00:00+00:00"}})
    assert b["files_with_baseline"] == 1 and b["first_baseline_at"].startswith("2026-10-01")


def test_field_list_from_another_format_of_the_same_data(tmp_root, monkeypatch):
    """IBAMA: the field list of 'Termo de apreensão - anexo' is only in its XML and JSON copies."""
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://p")
    desc = "Dicionário de dados:\n* SEQ_TAD – Chave.\n* NUM_TAD – Número.\n"
    pkg = {"name": "termo-de-apreensao", "title": "T", "resources": [
        {"id": "csv1", "name": "Termo de apreensão - anexo", "format": "CSV", "url": "https://p/a.csv", "description": ""},
        {"id": "xml1", "name": "Termo de apreensão - anexo", "format": "XML", "url": "https://p/a.xml", "description": desc},
        {"id": "csv2", "name": "Termo de apreensão - bem", "format": "CSV", "url": "https://p/b.csv", "description": ""}]}
    result, _ = census.run({}, packages=[pkg])
    t = {x["id"]: x for x in result["datasets"][0]["tables"]}
    assert t["csv1"]["level"] == 1 and t["csv1"]["link"]["method"] == "description-sibling"
    assert t["csv1"]["link"]["resource_id"] == "xml1"
    assert t["csv2"]["level"] == 0                                     # another table: not borrowed
    kind, schema = schemas.select("termo-de-apreensao", "csv1")
    assert kind == "described" and schema["x5ltep"]["source"]["resource_id"] == "xml1"


def test_history_row_and_file_memory(tmp_root):
    c = {"generated_at": "2026-10-12T03:30:00+00:00", "datasets": [{"name": "a", "tables": [
        {"id": "fixed"}, {"id": "broken"}, {"id": "same"}, {"id": "new"}]}]}
    v = {"fixed": {"pass_history": [["2026-10-05", False], ["2026-10-12", True]], "first_seen": "2026-10-01"},
         "broken": {"pass_history": [["2026-10-05", True], ["2026-10-12", False]], "first_seen": "2026-10-01"},
         "same": {"pass_history": [["2026-10-05", True]], "first_seen": "2026-10-01"},
         "new": {"pass_history": [["2026-10-12", True]], "first_seen": "2026-10-12T04:00:00+00:00"},
         "removed": {"pass_history": [], "first_seen": "2026-10-01"}}
    summary = {"generated_at": "x", "l1_rate": 0.5, "coverage": {"files": 4, "read": 4, "checked": 4, "conform": 3,
                                                               "not_downloaded": 0},
               "tables": {"by_level": {str(k): 0 for k in range(5)}}, "datasets": {"by_level": {str(k): 0 for k in range(5)}},
               "drift": {"observed": 1, "declared": 0}}
    row = report.history_row(c, v, summary, None)
    assert (row["fixed"], row["broken"], row["new_files"], row["gone_files"]) == (1, 1, 1, 1)
    rows = report.update_history(c, v, summary)
    rows = report.update_history(c, v, summary)                      # same chain reports again: replaced
    assert len(rows) == 1
    c2 = {**c, "generated_at": "2026-10-19T03:30:00+00:00"}
    rows = report.update_history(c2, v, summary)
    assert len(rows) == 2 and rows[-1]["gone_files"] == 0            # removed already counted last week


def test_network_figures_per_server():
    tables = [
        ({"url": "https://blob.example/a.zip"}, {"status": "ok", "kind": "zip", "bytes": 100e6, "seconds": 20,
                                                "timing": {"first_byte_s": 0.5, "download_s": 10}}),
        ({"url": "https://portal.example/b.csv"}, {"status": "ok", "kind": "csv", "bytes": 50e6, "seconds": 50,
                                                  "timing": {"first_byte_s": 4.0},
                                                  "network_failures": {"ConnectTimeout": 2}}),
        ({"url": "https://portal.example/c.csv"}, {"status": "error", "error": "HTTPError: 404 Client Error"}),
    ]
    n = report.network(tables)
    blob, portal = n["by_host"]["blob.example"], n["by_host"]["portal.example"]
    assert blob["zip_download_mb_s"] == 10.0 and blob["end_to_end_mb_s"] == 5.0
    assert portal["end_to_end_mb_s"] == 1.0 and portal["first_byte_s_median"] == 4.0
    assert portal["failed_files"] == 1 and portal["request_failures"] == {"ConnectTimeout": 2}
    assert n["total"]["files"] == 3 and n["total"]["gb"] == 0.15


def test_open_resource_records_timing(monkeypatch):
    serve(monkeypatch, {"u": _zip({"a.csv": b"x;y\n1;2\n"})})
    with tabular.open_resource("u") as src:
        for t in src.tables:
            list(t.rows)
    assert {"first_byte_s", "download_s"} <= set(src.digest["timing"])


def test_columns_without_name_are_a_finding_not_a_crash(tmp_root):
    """Recife, casos-de-dengue: 23 columns without a name in the header."""
    text = "a;;b;" + ";" * 0 + "\n1;2;3;4\n"
    out = validate.check_table(Table(text), None)
    names = [f["name"] for f in out["observed"]["fields"]]
    assert names == ["a", "_column_2", "b", "_column_4"]
    assert out["observed"]["fields"][1]["headerWithoutName"] is True
    assert validate.summarise([out], None)["columns_without_name"] == 2
    safety.check_schema({**out["observed"], "x5ltep": {"status": "observed"}}, {"observed"})


def test_a_suggested_schema_never_reaches_main(tmp_root):
    c, _ = _pdf_census()
    sneaky = {"fields": [{"name": "A"}], "x5ltep": {"status": "suggested"}}
    src = _artifact(tmp_root, {"schemas/ds/r1.extracted.json": sneaky})
    assert safety.accept_artifact(src, tmp_root / "repo", c, []) == []
    assert not (tmp_root / "repo/schemas/ds/r1.extracted.json").exists()


def test_one_malformed_schema_does_not_lose_the_batch(tmp_root):
    c, _ = _pdf_census()
    bad = {"fields": [{"name": ""}], "x5ltep": {"status": "observed"}}
    src = _artifact(tmp_root, {"results/validation.json": {"r1": {"status": "ok"}}, "schemas/ds/r1.observed.json": bad})
    accepted = safety.accept_artifact(src, tmp_root / "repo", c, [])
    assert accepted == ["results/validation.json"]               # the batch survives, the bad file is refused
    assert not (tmp_root / "repo/schemas/ds/r1.observed.json").exists()


def test_line_breaks_inside_header_names_are_cleaned():
    out = validate.observed_schema(["Origem\nDestino", "ok", "\t"], [["1", "2", "3"]])
    names = [f["name"] for f in out["fields"]]
    assert names == ["Origem Destino", "ok", "_column_3"]
    safety.check_schema({**out, "x5ltep": {"status": "observed"}}, {"observed"})


def test_similar_name_pairs_for_the_dashboard(tmp_root):
    c = {"datasets": [{"name": "ds", "title": "DS", "tables": [{"id": "r1", "name": "T"}, {"id": "r2", "name": "U"}]}]}
    v = {"r1": {"schema_kind": "declared", "summary": {"conformance": {"similar_names": {"IdeNuceloCEG": "IdeNucleoCEG"}}}},
         "r2": {"schema_kind": "extracted", "summary": {"conformance": {"probable_typos": {"NumCPFCNPJ": "NumCPF_CNPJ"}}}}}
    pairs = report.similar_name_pairs(c, v)                       # r2: record written before the rename
    assert {(p["declared"], p["column"]) for p in pairs} == {("NumCPFCNPJ", "NumCPF_CNPJ"), ("IdeNuceloCEG", "IdeNucleoCEG")}
    sims = [p["similarity"] for p in pairs]
    assert sims == sorted(sims, reverse=True) and all(0.8 <= x < 1.0 for x in sims)


# --- formats other than CSV ----------------------------------------------------------------

def _read_all(monkeypatch, data: bytes):
    serve(monkeypatch, {"u": data})
    out = []
    with tabular.open_resource("u") as src:
        for t in src.tables:
            out.append((t.fmt, t.member, list(t.rows)))
    return out, src.digest


def test_json_records_in_utf8_and_utf16(monkeypatch):
    recs = [{"A": 1, "B": "x"}, {"A": 2, "B": None, "C": [1, 2]}]
    out, digest = _read_all(monkeypatch, json.dumps(recs).encode("utf-8"))
    assert out == [("json", None, [["A", "B", "C"], ["1", "x", ""], ["2", "", "[1, 2]"]])]
    assert digest["kind"] == "json"
    out, _ = _read_all(monkeypatch, json.dumps({"dados": recs}, ensure_ascii=False).encode("utf-16"))
    assert out[0][2][0] == ["A", "B", "C"] and out[0][2][2][0] == "2"


def test_xml_records(monkeypatch):
    xml = b"<root><item><A>1</A><B>x</B></item><item><A>2</A><B>y</B></item></root>"
    out, _ = _read_all(monkeypatch, xml)
    assert out == [("xml", None, [["A", "B"], ["1", "x"], ["2", "y"]])]


def test_spreadsheets_and_parquet(monkeypatch):
    import openpyxl
    import pandas as pd
    import pyarrow as pa
    import pyarrow.parquet as pq

    wb = openpyxl.Workbook()
    wb.active.title = "Plan1"
    wb.active.append(["A", "B"])
    wb.active.append([1, 2.5])
    buf = io.BytesIO()
    wb.save(buf)
    out, digest = _read_all(monkeypatch, buf.getvalue())
    assert out == [("xlsx", "Plan1", [["A", "B"], ["1", "2.5"]])] and digest["kind"] == "xlsx"

    buf = io.BytesIO()
    pd.DataFrame({"A": ["1"], "B": ["x"]}).to_excel(buf, engine="odf", index=False, sheet_name="S")
    out, digest = _read_all(monkeypatch, buf.getvalue())
    assert out == [("ods", "S", [["A", "B"], ["1", "x"]])] and digest["kind"] == "ods"

    buf = io.BytesIO()
    pq.write_table(pa.table({"A": [1, 2], "B": ["x", None]}), buf)
    out, digest = _read_all(monkeypatch, buf.getvalue())
    assert out == [("parquet", None, [["A", "B"], ["1", "x"], ["2", ""]])] and digest["kind"] == "parquet"


def test_zip_inside_zip_and_zip_without_tables(monkeypatch):
    inner = _zip({"2007.csv": b"a;b\n1;2\n"})
    out, digest = _read_all(monkeypatch, _zip({"2007.zip": inner, "leia.pdf": b"%PDF-1"}))
    assert out == [("csv", "2007.zip/2007.csv", [["a", "b"], ["1", "2"]])] and digest["members"] == 1
    serve(monkeypatch, {"d": _zip({"doc.pdf": b"%PDF-1", "__MACOSX/._doc.pdf": b"x"})})
    with pytest.raises(tabular.NotTabular) as exc:
        with tabular.open_resource("d") as src:
            list(src.tables)
    assert exc.value.kind == "no-tables"


def _ods(sheets: dict[str, str]) -> bytes:
    """An ODS built by hand: content.xml with the rows given (already in ODS markup)."""
    ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')
    body = "".join(f'<table:table table:name="{n}">{rows}</table:table>' for n, rows in sheets.items())
    content = f'<office:document-content {ns}><office:body><office:spreadsheet>{body}</office:spreadsheet></office:body></office:document-content>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        zf.writestr("content.xml", content)
    return buf.getvalue()


def test_ods_is_read_row_by_row(monkeypatch):
    c = lambda v: f"<table:table-cell><text:p>{v}</text:p></table:table-cell>"
    empty = lambda n: f'<table:table-cell table:number-columns-repeated="{n}"/>'
    rows = ("<table:table-row>" + c("A") + c("B") + c("C") + empty(16000) + "</table:table-row>"
            + '<table:table-row><table:table-cell office:value-type="float" office:value="1.0"/>'
            + '<table:table-cell office:value-type="date" office:date-value="2022-03-01"/>'
            + '<table:table-cell office:value-type="boolean" office:boolean-value="TRUE"/></table:table-row>'
            + '<table:table-row table:number-rows-repeated="2">' + empty(5) + "</table:table-row>"
            + "<table:table-row>" + c("x") + "</table:table-row>"
            + '<table:table-row table:number-rows-repeated="1048000">' + empty(1024) + "</table:table-row>")
    serve(monkeypatch, {"u": _ods({"S1": rows, "S2": "<table:table-row>" + c("Z") + "</table:table-row>"})})
    with tabular.open_resource("u") as src:
        tables = iter(src.tables)
        first = next(tables)
        assert list(first.rows) == [["A", "B", "C"], ["1", "2022-03-01", "true"], ["", "", ""], ["", "", ""], ["x", "", ""]]
        second = next(tables)
        assert second.member == "S2" and list(second.rows) == [["Z"]]


def test_a_zip_larger_than_the_disk_is_read_as_it_streams(monkeypatch):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "P"
    wb.active.append(["X"])
    wb.active.append([7])
    xlsx = io.BytesIO()
    wb.save(xlsx)
    data = _zip({"a.csv": b"a;b\n1;2\n", "n.zip": _zip({"m.zip": _zip({"b.csv": b"c\n3\n"})}),
                 "p.xlsx": xlsx.getvalue(), "leia.pdf": b"%PDF-1"})
    monkeypatch.setattr(config, "MAX_ZIP_BYTES", 100)
    for headers in ({"Content-Length": str(len(data))}, {}):     # size known, or found out while downloading
        monkeypatch.setattr(tabular.ckan, "get", lambda url, **kw: FakeResponse(data, dict(headers)))
        with tabular.open_resource("u") as src:
            out = [(t.member, list(t.rows)) for t in src.tables]
        assert out == [("a.csv", [["a", "b"], ["1", "2"]]), ("n.zip/m.zip/b.csv", [["c"], ["3"]]),
                       ("p.xlsx#P", [["X"], ["7"]])]
        assert src.digest["streamed"] and src.digest["sha256"] == hashlib.sha256(data).hexdigest()
    # a spreadsheet larger than the disk cannot be read on this machine: said so, not dropped silently
    monkeypatch.setattr(tabular.ckan, "get", lambda url, **kw: FakeResponse(xlsx.getvalue()))
    with pytest.raises(tabular.NotTabular) as exc:
        with tabular.open_resource("u") as src:
            list(src.tables)
    assert exc.value.kind == "too-large"


def test_the_header_of_a_zip_costs_a_few_kilobytes_without_byte_ranges(monkeypatch):
    # Recife's file server answers a byte-range request with the whole file; the survey downloaded
    # 73 zips (21 GB) whole for their headers. Read from its start, a zip gives its first table's
    # header, and the download stops there.
    import random

    rnd = random.Random(1)
    body = "".join(f"{rnd.getrandbits(64):x};{rnd.getrandbits(64):x}\n" for _ in range(200_000)).encode()
    for compression in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", compression) as zf:
            zf.writestr("dados.csv", b"COD;VALOR\n" + body)
        data = buf.getvalue()
        resp = FakeResponse(data)
        monkeypatch.setattr(tabular.ckan, "get", lambda url, **kw: resp)
        assert tabular.peek_header("u", rows=0, stream_zip=True) == ("csv", ["COD", "VALOR"])
        assert len(data) > 3_000_000 and resp.raw.tell() < 300_000
    # a zip that cannot be read as a stream is downloaded whole
    monkeypatch.setattr(tabular.ckan, "get", lambda url, **kw: FakeResponse(data))
    monkeypatch.setattr(tabular, "_zip_stream_tables", lambda *a: (_ for _ in ()).throw(__import__("stream_unzip").TruncatedDataError()))
    assert tabular.peek_header("u", rows=0, stream_zip=True) == ("csv", ["COD", "VALOR"])


def test_zips_inside_zips_are_read_several_levels_deep(monkeypatch):
    out, _ = _read_all(monkeypatch, _zip({"1.zip": _zip({"2.zip": _zip({"3.csv": b"a\n1\n"})})}))
    assert out == [("csv", "1.zip/2.zip/3.csv", [["a"], ["1"]])]


def test_a_page_instead_of_the_table_is_a_link_failure(monkeypatch):
    serve(monkeypatch, {"h": b"<!DOCTYPE html><html>erro</html>"})
    with pytest.raises(tabular.NotTabular) as exc:
        with tabular.open_resource("h") as src:
            list(src.tables)
    assert exc.value.kind == "html"


def test_same_table_in_several_formats_is_one_table():
    mk = lambda i, name, fmt: {"id": i, "name": name, "candidate": fmt, "url": f"https://p/{i}"}
    tables = [mk("x", "Autos de infração", "xml"), mk("c", "Autos de infração", "csv"), mk("j", "Autos de infração", "json"),
              mk("a1", "auto-infracao.csv", "csv"), mk("a2", "auto-infracao.parquet", "parquet"),
              mk("y1", "Dados", "csv"), mk("y2", "Dados", "csv")]
    out = census.group_distributions(tables)
    by_id = {t["id"]: t for t in out}
    assert set(by_id) == {"c", "a1", "y1", "y2"}
    assert [d["format"] for d in by_id["c"]["distributions"]] == ["json", "xml"]
    assert [d["format"] for d in by_id["a1"]["distributions"]] == ["parquet"]
    assert "distributions" not in by_id["y1"]                         # two CSVs: two different files


def test_other_formats_are_checked_for_their_columns(monkeypatch):
    monkeypatch.setattr(work.ckan, "head", lambda url: {"content_length": 10})
    headers = {"https://p/j": ("json", ["A", "B"]), "https://p/x": ("xml", ["A", "C"])}
    monkeypatch.setattr(work.tabular, "peek_header", lambda url: headers[url])
    out = work.check_distributions([{"id": "j", "format": "json", "url": "https://p/j"},
                                    {"id": "x", "format": "xml", "url": "https://p/x"}], ["A", "B"])
    assert out[0]["same_columns"] is True
    assert out[1]["same_columns"] is False and out[1]["missing"] == ["b"] and out[1]["extra"] == ["c"]


def test_other_formats_differing_only_in_spelling_are_not_another_structure():
    # IBAMA: the CSV says "Ano Debito" and "Classe de Risco", the JSON says anoDebito and classeRisco
    d = {"status": "ok", "same_columns": False, "missing": ["ano debito", "classe de risco", "ano"],
         "extra": ["anodebito", "classe risco"]}
    out = report.with_spelling(d)
    assert out["spelling"] == {"ano debito": "anodebito", "classe de risco": "classe risco"}
    assert report.with_spelling({**d, "missing": ["numero de gru", "valor pago r"], "extra": ["numerogru", "valorpago"]})["spelling_only"]
    assert out["spelling_only"] is False and out["missing"] == ["ano"] and out["extra"] == []
    only = report.with_spelling({**d, "missing": ["ano debito"], "extra": ["anodebito"]})
    assert only["spelling_only"] is True
    s = report.distributions([({"distributions": [{}]}, {"distributions": [d, only]})])
    assert (s["spelling_only"], s["different_columns"]) == (1, 1)


def test_zips_without_tables_leave_the_tabular_universe(tmp_root):
    c = {"portal": {"portal_url": "https://p", "name": "P", "title": "P"}, "generated_at": "t", "datasets": [
        {"name": "a", "title": "A", "url": "u", "dictionaries": [], "tables": [
            {"id": "t1", "name": "T1", "url": "u", "level": 0, "link": None},
            {"id": "z1", "name": "Z1", "url": "u", "level": 0, "link": None}]}]}
    v = {"t1": {"status": "ok", "summary": {"rows": 1}},
         "z1": {"status": "not-tabular", "not_tabular_kind": "no-tables", "error": "zip without tables"}}
    s = report.build_summary(c, v, {}, [])
    assert s["tables"]["total"] == 1 and s["coverage"]["files"] == 1 and s["not_tables"]["count"] == 1


def test_a_dictionary_server_that_does_not_answer_is_asked_again_and_never_read_as_drift(tmp_root, monkeypatch):
    # Recife, 03/10/2026: the file server timed out for 69 dictionaries during one survey; their files
    # were checked against the DataStore instead, and a weekly run would have removed the declared
    # schemas (a false drift issue).
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://p")
    monkeypatch.setattr(census.config, "DICTIONARY_RETRY_WAIT_S", 0)
    serve(monkeypatch, {"https://p/d.csv": IBAMA_DICT})
    first, _ = census.run({}, packages=[_package()])
    calls = []

    def down(url, limit):
        calls.append(url)
        raise requests.ConnectTimeout("Max retries exceeded")

    monkeypatch.setattr(census.ckan, "fetch_bytes", down)
    second, events = census.run({}, packages=[_package()], previous=first)
    assert len(calls) == 1 + census.config.DICTIONARY_RETRY_ROUNDS
    ds = second["datasets"][0]
    assert ds["unreachable"] == ["r-dict"] and ds["kept_from"] == first["generated_at"]
    assert ds["tables"] == first["datasets"][0]["tables"]
    assert events == [] and schemas.select("termo-de-doacao", "r-data")[0] == "declared"
    # it answers on a retry: the fresh reading is used
    answers = iter([requests.ConnectTimeout("x")])

    def flaky(url, limit):
        e = next(answers, None)
        if e:
            raise e
        return IBAMA_DICT, hashlib.sha256(IBAMA_DICT).hexdigest()

    monkeypatch.setattr(census.ckan, "fetch_bytes", flaky)
    third, _ = census.run({}, packages=[_package()], previous=second)
    assert "kept_from" not in third["datasets"][0] and third["datasets"][0]["tables"][0]["level"] == 2
    # a 404 is the portal's broken link, not a server that did not answer
    resp = requests.Response()
    resp.status_code = 404
    assert census._download_error_kind(requests.HTTPError(response=resp)) == "download-failed"
    assert census._download_error_kind(requests.ConnectTimeout("x")) == "unreachable"


def test_a_survey_stops_asking_a_file_server_that_is_down_and_its_files_wait(tmp_root, monkeypatch):
    # Recife, 03/10/2026: the file server was down for the whole survey (almost five hours of retries).
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://p")
    monkeypatch.setattr(census.config, "DICTIONARY_RETRY_WAIT_S", 0)
    monkeypatch.setattr(census.config, "UNREACHABLE_STREAK", 2)
    monkeypatch.setattr(census.config, "CENSUS_WORKERS", 1)
    calls = []

    def down(url, limit):
        calls.append(url)
        raise requests.ConnectTimeout("Max retries exceeded")

    serve(monkeypatch, {})
    monkeypatch.setattr(census.ckan, "fetch_bytes", down)
    pkgs = [{**_package(dict_url=f"https://p/d{i}.csv"), "name": f"ds-{i}"} for i in range(5)]
    result, _ = census.run({}, packages=pkgs)
    # per round: two failures open the breaker, the other three datasets are not asked
    assert len(calls) == 2 * (1 + census.config.DICTIONARY_RETRY_ROUNDS)
    assert all(ds["unreachable"] == ["r-dict"] for ds in result["datasets"])
    assert result["datasets"][4]["dictionaries"][0]["error"] == census.NOT_ASKED
    # with no earlier reading, the files are not checked against a weaker schema: they wait
    assert work.plan(result, {}) == []


def test_a_header_the_server_did_not_give_is_asked_again(tmp_root, monkeypatch):
    # the header links a file to its dictionary: without it the link is not known
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://p")
    monkeypatch.setattr(census.config, "DICTIONARY_RETRY_WAIT_S", 0)
    serve(monkeypatch, {"https://p/d.csv": IBAMA_DICT, "https://p/x.pdf": b"%PDF-1.4 x"})
    answers = iter([requests.ConnectTimeout("x")])

    def peek(url, rows=0, **kw):
        e = next(answers, None)
        if e:
            raise e
        return None, ["SEQ_TAD", "DAT_TAD", "UF"]

    monkeypatch.setattr(census.tabular, "peek_header", peek)
    pkg = {"name": "termo-de-doacao", "title": "T", "resources": [
        {"id": "r-a", "name": "Planilha A", "format": "CSV", "url": "https://p/a.csv"},
        {"id": "r-dict", "name": "Metadados - Termo de doação", "format": "CSV", "url": "https://p/d.csv"},
        {"id": "r-pdf", "name": "Dicionário de outra coisa", "format": "PDF", "url": "https://p/x.pdf"}]}
    result, _ = census.run({}, packages=[pkg])
    t = result["datasets"][0]["tables"][0]
    assert t["link"]["method"] == "header" and "unreachable" not in result["datasets"][0]
    # never answering: the file waits
    monkeypatch.setattr(census.tabular, "peek_header", lambda url, rows=0, **kw: (_ for _ in ()).throw(requests.ConnectTimeout("x")))
    result, _ = census.run({}, packages=[pkg])
    assert result["datasets"][0]["unreachable"] == ["r-a"] and work.plan(result, {}) == []
    # with one dictionary there is no choice: the header does not decide the link, nothing waits
    pkg["resources"] = pkg["resources"][:2]
    result, _ = census.run({}, packages=[pkg])
    assert "unreachable" not in result["datasets"][0] and len(work.plan(result, {})) == 1


def test_a_header_is_read_in_the_survey_only_where_it_decides_something(tmp_root, monkeypatch):
    # Recife, 03/10/2026: the survey downloaded 73 zips (21 GB) whole for their headers, most of them
    # named by their dictionary's list of resource ids, where the header changes nothing.
    named = {"dictionary": "d", "part": 0, "method": "declared-id", "score": 1.0}
    single = {"dictionary": "d", "part": 0, "method": "single", "score": 0.5}
    assert not linker.header_decides([_d("d", "D", names=["a"]), _d("e", "E", names=["b"])], named)
    assert not linker.header_decides([_d("d", "Dicionário", names=["a"])], single)
    assert linker.header_decides([_d("d", "Metadados - Termos", names=["a"])], None)
    assert linker.header_decides([_d("d", "D", names=["a"]), _d("e", "E", names=["b"])], single)
    assert not linker.header_decides([_d("p", "Dicionário de dados")], None)       # a PDF lists no names
    # a PDF dictionary needs one header for its oracle: the cheapest file is read, and only it
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://p")
    serve(monkeypatch, {"https://p/x.pdf": b"%PDF-1.4 x"})
    asked = []
    monkeypatch.setattr(census.tabular, "peek_header", lambda url, rows=0, **kw: (asked.append(url), (None, ["A", "B"]))[1])
    pkg = {"name": "ds", "title": "T", "resources": [
        {"id": "r-zip", "name": "Dados 2020", "format": "ZIP", "url": "https://p/2020.zip", "size": 10},
        {"id": "r-big", "name": "Dados 2021", "format": "CSV", "url": "https://p/2021.csv", "size": 900},
        {"id": "r-small", "name": "Dados 2022", "format": "CSV", "url": "https://p/2022.csv", "size": 50},
        {"id": "r-pdf", "name": "Dicionário de dados", "format": "PDF", "url": "https://p/x.pdf"}]}
    result, _ = census.run({}, packages=[pkg])
    assert asked == ["https://p/2022.csv"]
    assert extract.tasks(result, {}, {}, "deterministic")[0]["header"] == ["A", "B"]
    # past SURVEY_HEADER_MAX_MINUTES no header is read: the file is linked without it and validated,
    # and the survey job is not stopped at its limit with all its work lost
    monkeypatch.setattr(census.config, "SURVEY_HEADER_MAX_MINUTES", -1)
    asked.clear()
    result, _ = census.run({}, packages=[pkg])
    assert asked == [] and any(t.get("header_deferred") for t in result["datasets"][0]["tables"])
    assert "unreachable" not in result["datasets"][0] and len(work.plan(result, {})) == 3


def test_waiting_for_a_server_is_bounded_and_a_redirect_loop_is_the_portals(tmp_root, monkeypatch):
    monkeypatch.setenv("CKAN_PORTAL_URL", "https://p")
    monkeypatch.setattr(census.config, "DICTIONARY_RETRY_WAIT_S", 0)
    serve(monkeypatch, {})
    monkeypatch.setattr(census.ckan, "fetch_bytes", lambda url, limit: (_ for _ in ()).throw(requests.ConnectTimeout("x")))
    first, _ = census.run({}, packages=[_package()])
    assert first["datasets"][0]["hold"] is True and work.plan(first, {}) == []
    # a file waits once: the next run uses what it observed
    second, _ = census.run({}, packages=[_package()], previous=first)
    assert "hold" not in second["datasets"][0] and [q["id"] for q in work.plan(second, {})] == ["r-data"]
    # a reading is kept for at most ROTATION_DAYS
    serve(monkeypatch, {"https://p/d.csv": IBAMA_DICT})
    good, _ = census.run({}, packages=[_package()])
    monkeypatch.setattr(census.ckan, "fetch_bytes", lambda url, limit: (_ for _ in ()).throw(requests.ConnectTimeout("x")))
    kept, _ = census.run({}, packages=[_package()], previous=good)
    assert kept["datasets"][0]["kept_from"] == good["generated_at"]
    old = {**kept, "datasets": [{**kept["datasets"][0], "kept_from": "2020-01-01T00:00:00+00:00"}]}
    late, _ = census.run({}, packages=[_package()], previous=old)
    assert "kept_from" not in late["datasets"][0]
    assert census._download_error_kind(requests.TooManyRedirects("Exceeded 30 redirects")) == "download-failed"
    assert census._download_error_kind(urllib3.exceptions.ProtocolError("IncompleteRead")) == "unreachable"
